"""End-to-end tests for persisted, evidence-backed pro-rata calculation runs."""

from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
import json
import uuid
from pathlib import Path

from openpyxl import load_workbook

from group_claims_backend.services.billing import BillingError, BillingService
from group_claims_backend.services.document_import import DocumentImportService
from group_claims_backend.services.reconciliation import ReconciliationService
from group_claims_backend.storage.service import StorageService


ROOT = Path(__file__).resolve().parents[2]
GENERATOR = ROOT / "scripts" / "validation-sprint"
SOURCES = (
    "01_opening_holdings_2025-01-31.xlsx",
    "02_partial_transfer_effective_2025-02-15.pdf",
    "03_claim_right_assignment_effective_2025-02-18.pdf",
    "04_later_holdings_2025-02-28.xlsx",
    "05_contradictory_custodian_statement_2025-02-28.pdf",
    "06_approved_billing_instruction_2025-03-01.pdf",
)


class BillingPipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.sources = self.root / "sources"
        subprocess.run([sys.executable, str(GENERATOR / "build_synthetic_workbooks.py"), str(self.sources)], check=True, capture_output=True)
        subprocess.run([sys.executable, str(GENERATOR / "build_synthetic_pdfs.py"), str(self.sources)], check=True, capture_output=True)
        self.storage = StorageService(str(self.root / "app-data"))
        self.storage.initialize()
        self.imports = DocumentImportService(self.storage)
        self.reconciliation = ReconciliationService(self.storage)
        self.billing = BillingService(self.storage)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def _prepare_complete_reconciliation(self) -> str:
        response = self.imports.import_authorized_paths([str(self.sources / item) for item in SOURCES])
        self.assertTrue(all(item["ok"] for item in response["documents"]))
        assert self.storage.database is not None
        with self.storage.database.connection() as connection:
            candidates = connection.execute("SELECT id, candidate_type FROM reconciliation_candidates ORDER BY created_at, id").fetchall()
            billing_candidate = connection.execute("SELECT id FROM billing_rule_candidates").fetchone()
        self.assertIsNotNone(billing_candidate)
        ordinary = [str(row["id"]) for row in candidates if row["candidate_type"] != "claim_assignment"]
        self.reconciliation.apply_policy({"candidateIds": ordinary, "policy": "reviewed_batch", "acknowledged": True})
        assignment = next(str(row["id"]) for row in candidates if row["candidate_type"] == "claim_assignment")
        self.reconciliation.confirm({"candidateId": assignment, "acknowledged": True})
        return str(billing_candidate["id"])

    def _approved_rule_with_payload(self, payload: dict[str, object]) -> str:
        """Insert an isolated deliberately malformed/edge approved rule.

        This models defensive handling of already-persisted bad data; normal
        rules always enter through candidate generation and confirmation.
        """
        assert self.storage.database is not None
        with self.storage.database.transaction() as connection:
            source = connection.execute("SELECT id FROM source_documents ORDER BY id LIMIT 1").fetchone()
            agreement_id = str(uuid.uuid4())
            connection.execute(
                "INSERT INTO billing_agreements (id,agreement_reference,effective_from,effective_to,terms_reference,created_at,source_document_id,source_reference,rule_json,approval_status,approved_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'approved', ?)",
                (agreement_id, "test-edge-rule", "2025-02-01", "2025-02-28", "test", "2025-03-01T00:00:00+00:00", source["id"], "page 1 block 0", json.dumps(payload), "2025-03-01T00:00:00+00:00"),
            )
        return agreement_id

    def _saved_allocation_ids(self, agreement_id: str) -> dict[str, str]:
        """Save a calculation solely to create immutable allocation IDs for an invoice test.

        The six-document fixture has no invoice.  These helpers add explicit
        verified ledger records in an isolated test database; they do not
        alter the fixture's ground truth or emulate document extraction.
        """
        run = self.billing.calculate({"billingAgreementId": agreement_id, "acknowledged": True})
        assert self.storage.database is not None
        with self.storage.database.connection() as connection:
            rows = connection.execute(
                "SELECT participant_key, id FROM billing_calculation_allocations WHERE calculation_run_id = ?",
                (run["id"],),
            ).fetchall()
        return {str(row["participant_key"]): str(row["id"]) for row in rows}

    def _verified_invoice_line(
        self, agreement_id: str, allocation_id: str, amount: str, *, status: str = "issued",
        currency: str = "USD", payment: str | None = None, credit: str | None = None,
    ) -> tuple[str, str]:
        """Seed an independently verified invoice ledger record for one test.

        Production imports and AI output do not call this helper.  The source
        document and evidence references make the test record eligible for a
        read-only comparison only after its explicit issued lifecycle state.
        """
        assert self.storage.database is not None
        now = "2025-03-05T00:00:00+00:00"
        source_document_id, entity_id, recipient_id, invoice_id, line_id = (str(uuid.uuid4()) for _ in range(5))
        with self.storage.database.transaction() as connection:
            connection.execute(
                "INSERT INTO source_documents (id,original_filename,managed_relative_path,content_sha256,file_size_bytes,imported_at,file_format,processing_status) "
                "VALUES (?, 'verified-invoice-fixture.pdf', ?, ?, 1, ?, 'pdf', 'parsed')",
                (source_document_id, f"documents/{source_document_id}", (source_document_id.replace("-", "") * 2)[:64], now),
            )
            connection.execute(
                "INSERT INTO legal_entities (id,display_name,legal_name,entity_type,jurisdiction,created_at) VALUES (?, ?, NULL, 'company', NULL, ?)",
                (entity_id, "External Billing Agent", now),
            )
            connection.execute(
                "INSERT INTO invoice_recipients (id,legal_entity_id,created_at) VALUES (?, ?, ?)",
                (recipient_id, entity_id, now),
            )
            connection.execute(
                "INSERT INTO invoices (id,invoice_recipient_id,billing_agreement_id,status,issued_at,created_at) VALUES (?, ?, ?, ?, ?, ?)",
                (invoice_id, recipient_id, agreement_id, status, now if status != "draft" else None, now),
            )
            connection.execute(
                "INSERT INTO invoice_line_items (id,invoice_id,description,amount_decimal,currency_code,created_at,billing_calculation_allocation_id) "
                "VALUES (?, ?, 'Verified fixture invoice line', ?, ?, ?, ?)",
                (line_id, invoice_id, amount, currency, now, allocation_id),
            )
            connection.execute(
                "INSERT INTO evidence_references VALUES (?, ?, 'invoice_line_item', ?, 'page 1 block 0', 'Verified invoice-line evidence', ?)",
                (str(uuid.uuid4()), source_document_id, line_id, now),
            )
            if payment is not None:
                payment_id = str(uuid.uuid4())
                connection.execute(
                    "INSERT INTO payments (id,invoice_id,payer_entity_id,amount_decimal,currency_code,received_at,created_at) VALUES (?, ?, NULL, ?, ?, ?, ?)",
                    (payment_id, invoice_id, payment, currency, now, now),
                )
                connection.execute(
                    "INSERT INTO evidence_references VALUES (?, ?, 'payment', ?, 'page 1 block 0', 'Verified payment evidence', ?)",
                    (str(uuid.uuid4()), source_document_id, payment_id, now),
                )
            if credit is not None:
                adjustment_id = str(uuid.uuid4())
                connection.execute(
                    "INSERT INTO financial_adjustments (id,invoice_id,adjustment_type,amount_decimal,currency_code,reason,created_at) VALUES (?, ?, 'credit_note', ?, ?, 'Verified credit note', ?)",
                    (adjustment_id, invoice_id, credit, currency, now),
                )
                connection.execute(
                    "INSERT INTO evidence_references VALUES (?, ?, 'financial_adjustment', ?, 'page 1 block 0', 'Verified credit-note evidence', ?)",
                    (str(uuid.uuid4()), source_document_id, adjustment_id, now),
            )
        return invoice_id, line_id

    def test_approved_rule_calculates_persisted_provisional_allocation_and_survives_restart(self) -> None:
        candidate_id = self._prepare_complete_reconciliation()
        candidate = self.billing.candidate({"candidateId": candidate_id})
        self.assertEqual(candidate["status"], "candidate")
        self.assertEqual(candidate["payload"]["feePool"], "120000.00")
        rule = self.billing.approve_rule({"candidateId": candidate_id, "acknowledged": True})
        self.assertEqual(rule["status"], "approved")
        self.assertEqual(rule["invoiceRecipientStatus"], "unresolved")

        preview = self.billing.preview({"billingAgreementId": rule["id"]})
        self.assertEqual(preview["status"], "provisional")
        self.assertEqual(preview["feePool"], "120000.00")
        self.assertEqual(preview["totalEligibleAmount"], "1400000")
        allocations = {item["subfundReference"]: item for item in preview["allocations"]}
        self.assertEqual(allocations["SF-A"]["allocatedFee"], "34285.71")
        self.assertEqual(allocations["SF-B"]["allocatedFee"], "85714.29")
        self.assertEqual(allocations["SF-C"]["eligibilityStatus"], "excluded")
        self.assertIsNone(allocations["SF-C"]["allocatedFee"])
        self.assertEqual(preview["allocatedTotal"], "120000.00")
        self.assertEqual(len(preview["relatedIssues"]), 1)
        self.assertTrue(preview["relatedIssues"][0]["sourceReference"])
        self.assertEqual(allocations["SF-A"]["issuedInvoiceStatus"], "no_issued_invoice_recorded")
        self.assertIsNone(allocations["SF-A"]["verifiedInvoicedAmount"])
        self.assertIsNone(allocations["SF-A"]["outstandingBalance"])

        saved = self.billing.calculate({"billingAgreementId": rule["id"], "acknowledged": True})
        self.assertEqual(saved["version"], 1)
        self.assertEqual(saved["status"], "provisional")
        repeated = self.billing.calculate({"billingAgreementId": rule["id"], "acknowledged": True})
        self.assertEqual(repeated["id"], saved["id"])
        self.assertEqual(len(self.billing.runs({"billingAgreementId": rule["id"]})["runs"]), 1)

        restarted = StorageService(str(self.root / "app-data"))
        restarted.initialize()
        restored = BillingService(restarted).run({"runId": saved["id"]})
        self.assertEqual(restored["feePool"], "120000.00")
        self.assertEqual({item["subfundReference"]: item["allocatedFee"] for item in restored["allocations"]}["SF-B"], "85714.29")

    def test_verified_issued_invoice_payment_and_credit_are_separate_from_allocation(self) -> None:
        candidate_id = self._prepare_complete_reconciliation()
        rule = self.billing.approve_rule({"candidateId": candidate_id, "acknowledged": True})
        allocation_ids = self._saved_allocation_ids(str(rule["id"]))
        self._verified_invoice_line(
            str(rule["id"]), allocation_ids["SF-A-1"], "34285.71", status="partially_paid",
            payment="1000.00", credit="200.00",
        )

        preview = self.billing.preview({"billingAgreementId": rule["id"]})
        allocations = {item["subfundReference"]: item for item in preview["allocations"]}
        fund_a = allocations["SF-A"]
        self.assertEqual(fund_a["allocatedFee"], "34285.71")
        self.assertEqual(fund_a["issuedInvoiceStatus"], "verified_issued")
        self.assertEqual(fund_a["verifiedInvoicedAmount"], "34285.71")
        self.assertEqual(fund_a["allocationInvoiceDifference"], "0.00")
        self.assertEqual(fund_a["outstandingBalance"], "33085.71")
        self.assertEqual(fund_a["outstandingStatus"], "verified")
        self.assertEqual(fund_a["invoiceRecipient"], "External Billing Agent")
        self.assertTrue(fund_a["invoiceEvidence"])

    def test_draft_unpaid_and_mixed_currency_invoices_are_not_presented_as_verified_amounts(self) -> None:
        candidate_id = self._prepare_complete_reconciliation()
        rule = self.billing.approve_rule({"candidateId": candidate_id, "acknowledged": True})
        allocation_ids = self._saved_allocation_ids(str(rule["id"]))

        self._verified_invoice_line(str(rule["id"]), allocation_ids["SF-B-2"], "85714.29", status="draft")
        draft_preview = self.billing.preview({"billingAgreementId": rule["id"]})
        draft_b = {item["subfundReference"]: item for item in draft_preview["allocations"]}["SF-B"]
        self.assertEqual(draft_b["issuedInvoiceStatus"], "draft_only")
        self.assertIsNone(draft_b["verifiedInvoicedAmount"])

        # A later explicit issued record in another currency must not be
        # combined with the USD allocation or shown as a comparable amount.
        self._verified_invoice_line(str(rule["id"]), allocation_ids["SF-B-2"], "75000.00", currency="EUR")
        currency_preview = self.billing.preview({"billingAgreementId": rule["id"]})
        currency_b = {item["subfundReference"]: item for item in currency_preview["allocations"]}["SF-B"]
        self.assertEqual(currency_b["issuedInvoiceStatus"], "currency_mismatch")
        self.assertIsNone(currency_b["verifiedInvoicedAmount"])
        self.assertIsNone(currency_b["allocationInvoiceDifference"])
        self.assertIsNone(currency_b["outstandingBalance"])

    def test_issued_amount_can_differ_from_calculation_but_needs_payment_evidence_for_outstanding(self) -> None:
        candidate_id = self._prepare_complete_reconciliation()
        rule = self.billing.approve_rule({"candidateId": candidate_id, "acknowledged": True})
        allocation_ids = self._saved_allocation_ids(str(rule["id"]))
        self._verified_invoice_line(str(rule["id"]), allocation_ids["SF-B-2"], "80000.00")

        preview = self.billing.preview({"billingAgreementId": rule["id"]})
        fund_b = {item["subfundReference"]: item for item in preview["allocations"]}["SF-B"]
        self.assertEqual(fund_b["issuedInvoiceStatus"], "verified_issued")
        self.assertEqual(fund_b["verifiedInvoicedAmount"], "80000.00")
        self.assertEqual(fund_b["allocationInvoiceDifference"], "5714.29")
        self.assertIsNone(fund_b["outstandingBalance"])
        self.assertEqual(fund_b["outstandingStatus"], "payment_information_unavailable")

    def test_invoice_line_linked_to_another_agreement_is_not_misattributed(self) -> None:
        candidate_id = self._prepare_complete_reconciliation()
        rule = self.billing.approve_rule({"candidateId": candidate_id, "acknowledged": True})
        candidate = self.billing.candidate({"candidateId": candidate_id})
        other_rule = self._approved_rule_with_payload(dict(candidate["payload"]))
        other_allocations = self._saved_allocation_ids(other_rule)

        # The invoice belongs to the first agreement but its line deliberately
        # references a run for another agreement.  It must remain unmatched,
        # rather than being assigned by description, subfund name, or amount.
        self._verified_invoice_line(str(rule["id"]), other_allocations["SF-A-1"], "34285.71")
        preview = self.billing.preview({"billingAgreementId": rule["id"]})
        fund_a = {item["subfundReference"]: item for item in preview["allocations"]}["SF-A"]
        self.assertEqual(fund_a["issuedInvoiceStatus"], "issued_unmapped")
        self.assertIsNone(fund_a["verifiedInvoicedAmount"])
        self.assertEqual(preview["invoicing"]["unmappedIssuedInvoiceLineCount"], 1)

    def test_missing_approval_and_invalid_calculation_requests_are_blocked(self) -> None:
        candidate_id = self._prepare_complete_reconciliation()
        with self.assertRaisesRegex(BillingError, "Approved billing rule was not found"):
            self.billing.preview({"billingAgreementId": candidate_id})
        rule = self.billing.approve_rule({"candidateId": candidate_id, "acknowledged": True})
        with self.assertRaisesRegex(BillingError, "Explicit acknowledgement"):
            self.billing.calculate({"billingAgreementId": rule["id"]})
        with self.assertRaisesRegex(BillingError, "Dates must use YYYY-MM-DD"):
            self.billing.preview({"billingAgreementId": rule["id"], "asOfDate": "28/02/2025"})

    def test_billing_rule_candidate_is_idempotent_and_does_not_create_invoice_or_recipient(self) -> None:
        candidate_id = self._prepare_complete_reconciliation()
        assert self.storage.database is not None
        with self.storage.database.connection() as connection:
            document = connection.execute("SELECT source_document_id FROM billing_rule_candidates WHERE id = ?", (candidate_id,)).fetchone()
        regenerated = self.billing.generate({"documentId": str(document["source_document_id"])})
        self.assertEqual(regenerated["candidates"][0]["id"], candidate_id)
        first = self.billing.approve_rule({"candidateId": candidate_id, "acknowledged": True})
        repeated = self.billing.approve_rule({"candidateId": candidate_id, "acknowledged": True})
        self.assertEqual(repeated["id"], first["id"])
        with self.storage.database.connection() as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM invoices").fetchone()[0], 0)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM invoice_recipients").fetchone()[0], 0)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM billing_agreements WHERE billing_rule_candidate_id = ?", (candidate_id,)).fetchone()[0], 1)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM audit_events WHERE event_type = 'billing_rule_approved' AND target_id = ?", (first["id"],)).fetchone()[0], 1)

    def test_billing_rule_review_queue_exposes_only_local_provisional_rules(self) -> None:
        candidate_id = self._prepare_complete_reconciliation()
        queue = self.billing.candidates({})["candidates"]
        self.assertEqual([item["id"] for item in queue], [candidate_id])
        self.assertEqual(queue[0]["status"], "candidate")
        self.assertIsNone(queue[0]["billingAgreementId"])
        self.assertTrue(queue[0]["sourceDocumentId"])
        rule = self.billing.approve_rule({"candidateId": candidate_id, "acknowledged": True})
        self.assertEqual(self.billing.candidates({"status": "candidate"})["candidates"], [])
        approved = self.billing.candidates({"status": "approved"})["candidates"]
        self.assertEqual(len(approved), 1)
        self.assertEqual(approved[0]["id"], candidate_id)
        self.assertEqual(approved[0]["billingAgreementId"], rule["id"])
        self.assertNotEqual(approved[0]["id"], approved[0]["billingAgreementId"])

    def test_approved_candidate_without_linked_agreement_is_reported_without_identifier_fallback(self) -> None:
        candidate_id = self._prepare_complete_reconciliation()
        assert self.storage.database is not None
        with self.storage.database.transaction() as connection:
            connection.execute("UPDATE billing_rule_candidates SET lifecycle_status = 'approved' WHERE id = ?", (candidate_id,))
        approved = self.billing.candidates({"status": "approved"})["candidates"]
        self.assertEqual(approved[0]["id"], candidate_id)
        self.assertIsNone(approved[0]["billingAgreementId"])

    def test_recalculation_versions_inputs_without_replacing_an_earlier_run(self) -> None:
        candidate_id = self._prepare_complete_reconciliation()
        rule = self.billing.approve_rule({"candidateId": candidate_id, "acknowledged": True})
        first = self.billing.calculate({"billingAgreementId": rule["id"], "acknowledged": True})

        # This later synthetic statement changes the reconciled record linked
        # to a calculation request.  It cannot retroactively alter the
        # explicitly approved 1,000,000 billing basis or replace run one.
        later = self.sources / "later-holdings-2025-03-01.xlsx"
        workbook = load_workbook(self.sources / "04_later_holdings_2025-02-28.xlsx")
        sheet = workbook["Holdings"]
        sheet["B3"] = "2025-03-01"
        sheet["B4"] = "2025-03-02"
        for row in range(8, 11):
            sheet.cell(row, 8).value = "2025-03-01"
        sheet["G9"] = 1_100_000
        workbook.save(later)
        imported = self.imports.import_authorized_paths([str(later)])
        self.assertTrue(imported["documents"][0]["ok"])
        assert self.storage.database is not None
        with self.storage.database.connection() as connection:
            rows = connection.execute(
                "SELECT c.id FROM reconciliation_candidates c JOIN source_documents d ON d.id = c.source_document_id WHERE d.original_filename = ?",
                (later.name,),
            ).fetchall()
        self.reconciliation.apply_policy({"candidateIds": [str(row["id"]) for row in rows], "policy": "reviewed_batch", "acknowledged": True})
        second = self.billing.calculate({"billingAgreementId": rule["id"], "asOfDate": "2025-03-01", "acknowledged": True})
        self.assertEqual(second["version"], 2)
        self.assertEqual(second["previousRunId"], first["id"])
        self.assertEqual({item["subfundReference"]: item["allocatedFee"] for item in second["allocations"]}["SF-B"], "85714.29")
        comparison = self.billing.compare_runs({"leftRunId": first["id"], "rightRunId": second["id"]})
        self.assertEqual(comparison["billingAgreementId"], rule["id"])
        restored = self.billing.run({"runId": first["id"]})
        self.assertEqual(restored["version"], 1)

    def test_invalid_dates_zero_bases_and_rounding_residuals_are_deterministic(self) -> None:
        self._prepare_complete_reconciliation()
        base = {
            "periodStart": "2025-02-01", "periodEnd": "2025-02-28", "allocationDate": None,
            "allocationDateRequired": True, "feePool": "1.00", "currency": "USD", "participants": [
                {"key": "a", "name": "Horizon Credit Fund A (SF-A)", "label": "A", "subfundReference": "SF-A", "eligibleBasis": "1", "status": "eligible", "reason": "test"},
            ],
        }
        missing_date = self._approved_rule_with_payload(base)
        with self.assertRaisesRegex(BillingError, "requires an allocation date"):
            self.billing.preview({"billingAgreementId": missing_date})
        zero = dict(base, allocationDateRequired=False, participants=[dict(base["participants"][0], eligibleBasis="0")])
        zero_rule = self._approved_rule_with_payload(zero)
        with self.assertRaisesRegex(BillingError, "Eligible nominal amounts must be positive"):
            self.billing.preview({"billingAgreementId": zero_rule})
        rounded = dict(base, allocationDateRequired=False, feePool="0.01", participants=[
            {"key": key, "name": "unmatched", "label": key, "subfundReference": key, "eligibleBasis": "1", "status": "eligible", "reason": "test"}
            for key in ("a", "b", "c")
        ])
        rounded_rule = self._approved_rule_with_payload(rounded)
        preview = self.billing.preview({"billingAgreementId": rounded_rule})
        self.assertEqual([item["allocatedFee"] for item in preview["allocations"]], ["0.01", "0.00", "0.00"])
        self.assertEqual(preview["allocatedTotal"], "0.01")

