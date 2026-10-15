"""Synthetic regression tests for deterministic spreadsheet extraction."""

from __future__ import annotations

import tempfile
import unittest
import json
import uuid
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from group_claims_backend.services.document_import import DeterministicDocumentParser, ExtractionUnsupported
from group_claims_backend.storage.service import StorageService
from group_claims_backend.services.document_import import DocumentImportService


@unittest.skipUnless(
    __import__("importlib").util.find_spec("openpyxl") is not None,
    "openpyxl is required for XLSX extraction tests",
)
class XlsxExtractionTests(unittest.TestCase):
    def test_extensionless_vault_copy_preserves_cells_and_provenance(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            original = Path(directory) / "source.xlsx"
            vault_copy = Path(directory) / "3b4fbb46-0b30-4e3b-8dcd-f2a67cdb2e03"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Holdings data"
            sheet["A1"] = "Account"
            sheet["A2"] = "000123"
            # Identifiers beyond Excel's numeric precision must be stored as text.
            sheet["B2"] = "12345678901234567890"
            sheet["C2"] = "=B2"
            sheet.merge_cells("D1:E1")
            sheet["D1"] = "Merged heading"
            workbook.create_sheet("Notes")["B4"] = "Irregular header position"
            workbook.save(original)
            vault_copy.write_bytes(original.read_bytes())

            items = DeterministicDocumentParser()._parse_xlsx(vault_copy)

        by_reference = {item.source_reference: item for item in items}
        self.assertEqual(by_reference["Holdings data!A2"].value_text, "000123")
        self.assertEqual(by_reference["Holdings data!B2"].value_text, "12345678901234567890")
        self.assertEqual(by_reference["Holdings data!C2"].formula_text, "=B2")
        self.assertTrue(by_reference["Holdings data!E1"].is_empty)
        self.assertEqual(by_reference["Notes!B4"].value_text, "Irregular header position")
        self.assertEqual(by_reference["Holdings data!A2"].row_number, 2)
        self.assertEqual(by_reference["Holdings data!A2"].column_number, 1)

    def test_invalid_workbook_has_controlled_error(self) -> None:
        from group_claims_backend.services.document_import import DocumentImportError

        with tempfile.TemporaryDirectory() as directory:
            invalid = Path(directory) / "vault-id"
            invalid.write_bytes(b"not an xlsx")
            with self.assertRaisesRegex(DocumentImportError, "Excel workbook could not be read"):
                DeterministicDocumentParser()._parse_xlsx(invalid)

    def test_parser_preserves_unicode_source_filename(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "Bestätigung März.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "test"
            workbook.save(source)
            storage = StorageService(str(Path(directory) / "app-data"))
            storage.initialize()
            result = DocumentImportService(storage).import_authorized_paths([str(source)])

        self.assertEqual(result["documents"][0]["document"]["original_filename"], "Bestätigung März.xlsx")

    def test_reimport_retries_a_previously_failed_duplicate(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "retry.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "Recovered"
            workbook.save(source)
            storage = StorageService(str(Path(directory) / "app-data"))
            storage.initialize()
            service = DocumentImportService(storage)
            first = service.import_authorized_paths([str(source)])["documents"][0]["document"]
            assert storage.database
            with storage.database.transaction() as connection:
                connection.execute(
                    "UPDATE source_documents SET processing_status = 'failed', processing_error = 'Earlier parser failure' WHERE id = ?",
                    (first["id"],),
                )

            retry = service.import_authorized_paths([str(source)])["documents"][0]

        self.assertTrue(retry["ok"])
        self.assertTrue(retry["duplicate"])
        self.assertEqual(retry["document"]["id"], first["id"])
        self.assertEqual(retry["document"]["processing_status"], "parsed")

    def test_xlsx_without_dimension_metadata_is_imported(self) -> None:
        """Workbooks created by some tools omit the optional dimension XML."""
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "without-dimension.xlsx"
            workbook = Workbook()
            workbook.active.title = "Holdings"
            workbook.active["A1"] = "Custodian account"
            workbook.active["B2"] = "AC-1001"
            workbook.save(source)

            rewritten = Path(directory) / "without-dimension-rewritten.xlsx"
            with zipfile.ZipFile(source) as original, zipfile.ZipFile(rewritten, "w") as target:
                for member in original.infolist():
                    payload = original.read(member.filename)
                    if member.filename == "xl/worksheets/sheet1.xml":
                        payload = payload.replace(b'<dimension ref="A1:B2"/>', b"")
                    target.writestr(member, payload)

            items = DeterministicDocumentParser()._parse_xlsx(rewritten)

        values = {item.source_reference: item.value_text for item in items}
        self.assertEqual(values["Holdings!A1"], "Custodian account")
        self.assertEqual(values["Holdings!B2"], "AC-1001")


class DocumentSummaryStatusTests(unittest.TestCase):
    def test_summary_marks_only_the_current_extraction_as_ai_reviewed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "holdings.csv"
            source.write_text("Fund,Quantity\nFund A,10\n", encoding="utf-8")
            storage = StorageService(str(Path(directory) / "app-data"))
            storage.initialize()
            service = DocumentImportService(storage)
            imported = service.import_authorized_paths([str(source)])["documents"][0]["document"]
            self.assertEqual(imported["ai_review_status"], "new")
            assert storage.database
            with storage.database.transaction() as connection:
                extraction = connection.execute(
                    "SELECT id FROM document_extractions WHERE source_document_id = ? AND status = 'parsed'",
                    (imported["id"],),
                ).fetchone()
                assert extraction is not None
                connection.execute(
                    "INSERT INTO ai_processing_jobs (id, source_document_id, extraction_id, provider, model, prompt_version, schema_version, status, created_at, completed_at) "
                    "VALUES (?, ?, ?, 'mock', 'test', 'v1', 'v1', 'completed', ?, ?)",
                    (str(uuid.uuid4()), imported["id"], extraction["id"], "2026-01-01T00:00:00+00:00", "2026-01-01T00:00:00+00:00"),
                )
            self.assertEqual(service.list_documents()["documents"][0]["ai_review_status"], "reviewed")

    def test_remove_document_deletes_only_managed_copy_and_derived_records(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "remove-me.csv"
            original_bytes = b"Fund,Quantity\nFund A,10\n"
            source.write_bytes(original_bytes)
            storage = StorageService(str(Path(directory) / "app-data"))
            storage.initialize()
            service = DocumentImportService(storage)
            imported = service.import_authorized_paths([str(source)])["documents"][0]["document"]
            assert storage.paths is not None
            managed = storage.paths.documents / imported["id"]
            self.assertTrue(managed.exists())

            removed = service.remove_document(imported["id"])

            self.assertEqual(removed["documentId"], imported["id"])
            self.assertEqual(source.read_bytes(), original_bytes)
            self.assertFalse(managed.exists())
            self.assertEqual(service.list_documents()["documents"], [])
            assert storage.database is not None
            with storage.database.connection() as connection:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM document_extractions").fetchone()[0], 0)
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM extracted_content").fetchone()[0], 0)


class SpreadsheetContentTests(unittest.TestCase):
    def test_grid_is_ordered_and_keeps_formula_and_sparse_positions(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            workbook = Workbook()
            data = workbook.active
            data.title = "Data"
            data["A1"] = "000123"
            data["C2"] = "late cell"
            data["B3"] = "=1+1"
            notes = workbook.create_sheet("Notes")
            notes["B4"] = "second worksheet"
            workbook.save(source)
            storage = StorageService(str(Path(directory) / "app-data"))
            storage.initialize()
            service = DocumentImportService(storage)
            imported = service.import_authorized_paths([str(source)])["documents"][0]["document"]
            content = service.content(imported["id"], 0, 25, "Data", 1, 3)

        self.assertEqual([sheet["name"] for sheet in content["worksheets"]], ["Data", "Notes"])
        self.assertEqual(content["grid"]["rowStart"], 1)
        cells = content["grid"]["cells"]
        self.assertEqual([(cell["row_number"], cell["column_number"]) for cell in cells], [(1, 1), (2, 3), (3, 2)])
        self.assertEqual(cells[0]["value_text"], "000123")
        self.assertEqual(cells[2]["formula_text"], "=1+1")
        self.assertEqual(content["extraction"]["nonEmptyItems"], 4)

    def test_grid_is_bounded_without_losing_coordinate_navigation(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "dense.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Dense"
            for row in range(1, 16):
                for column in range(1, 16):
                    sheet.cell(row=row, column=column, value=f"cell-{row}-{column}-{'x' * 200}")
            workbook.save(source)
            storage = StorageService(str(Path(directory) / "app-data"))
            storage.initialize()
            service = DocumentImportService(storage)
            imported = service.import_authorized_paths([str(source)])["documents"][0]["document"]
            first = service.content(imported["id"], 0, 1, "Dense", 1, 10, 1, 10)
            second = service.content(imported["id"], 0, 1, "Dense", 11, 10, 11, 10)

        self.assertEqual(len(first["grid"]["cells"]), 100)
        self.assertEqual(first["grid"]["cells"][0]["source_reference"], "Dense!A1")
        self.assertEqual(second["grid"]["cells"][0]["source_reference"], "Dense!K11")
        self.assertFalse(first["grid"]["previewTruncated"])
        self.assertLess(len(json.dumps(first, ensure_ascii=False).encode("utf-8")), 60_000)
        self.assertTrue(first["grid"]["cells"][0]["value_truncated"])


@unittest.skipUnless(
    __import__("importlib").util.find_spec("fitz") is not None,
    "PyMuPDF is required for PDF extraction tests",
)
class PdfExtractionTests(unittest.TestCase):
    def _image_bytes(self) -> bytes:
        from PIL import Image, ImageDraw, ImageFont

        image = Image.new("RGB", (1200, 300), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.load_default(size=48)
        draw.text((40, 100), "OCR sample 00123", fill="black", font=font)
        output = BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()

    def _mixed_pdf(self, target: Path) -> None:
        import fitz

        document = fitz.open()
        native = document.new_page()
        native.insert_text((72, 72), "Native PDF text")
        scanned = document.new_page(width=600, height=300)
        scanned.insert_image(fitz.Rect(0, 0, 600, 300), stream=self._image_bytes())
        document.save(target)
        document.close()

    def test_native_text_uses_coordinate_aware_blocks(self) -> None:
        import fitz

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "native.pdf"
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 72), "Native document text")
            document.save(path)
            document.close()

            result = DeterministicDocumentParser().parse(path, "pdf")

        self.assertEqual(result.pdf_pages[0].extraction_method, "native_text")
        self.assertIn("Native document text", result.pdf_pages[0].cleaned_text or "")
        self.assertIsNotNone(result.pdf_pages[0].blocks[0].x0)
        self.assertEqual(result.items[0].page_number, 1)

    @unittest.skipUnless(
        __import__("importlib").util.find_spec("pytesseract") is not None,
        "pytesseract is required for OCR tests",
    )
    def test_mixed_text_and_scanned_pages_ocr_only_missing_text_layer(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mixed.pdf"
            self._mixed_pdf(path)
            result = DeterministicDocumentParser().parse(path, "pdf")

        self.assertEqual(result.pdf_pages[0].extraction_method, "native_text")
        self.assertEqual(result.pdf_pages[1].extraction_method, "ocr")
        self.assertIn("Native PDF text", result.pdf_pages[0].cleaned_text or "")
        self.assertTrue((result.pdf_pages[1].cleaned_text or "").strip())
        self.assertIsNotNone(result.pdf_pages[1].quality_warning)

    def test_missing_ocr_runtime_is_a_controlled_unsupported_result(self) -> None:
        import pytesseract

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scanned.pdf"
            import fitz
            document = fitz.open()
            page = document.new_page(width=600, height=300)
            page.insert_image(fitz.Rect(0, 0, 600, 300), stream=self._image_bytes())
            document.save(path)
            document.close()
            with patch.object(pytesseract, "get_tesseract_version", side_effect=pytesseract.TesseractNotFoundError()):
                with self.assertRaisesRegex(ExtractionUnsupported, "OCR is unavailable"):
                    DeterministicDocumentParser().parse(path, "pdf")

    def test_pdf_page_provenance_is_persisted_for_a_new_extraction_attempt(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "app-data"
            source = Path(directory) / "native.pdf"
            import fitz
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 72), "Persisted PDF page")
            document.save(source)
            document.close()
            storage = StorageService(str(root))
            self.assertEqual(storage.initialize().status, "ready")
            response = DocumentImportService(storage).import_authorized_paths([str(source)])
            imported = response["documents"][0]
            self.assertTrue(imported["ok"])
            document_id = imported["document"]["id"]
            assert storage.database is not None
            with storage.database.connection() as connection:
                row = connection.execute("""SELECT p.page_number, p.extraction_method, b.text
                    FROM pdf_page_extractions p JOIN pdf_text_blocks b ON b.pdf_page_extraction_id = p.id
                    JOIN document_extractions e ON e.id = p.extraction_id
                    WHERE e.source_document_id = ?""", (document_id,)).fetchone()
            self.assertEqual(row["page_number"], 1)
            self.assertEqual(row["extraction_method"], "native_text")
            self.assertIn("Persisted PDF page", row["text"])

    def test_pdf_content_uses_source_page_numbers_for_navigation(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "five-pages.pdf"
            import fitz
            document = fitz.open()
            for page_number in range(1, 6):
                page = document.new_page()
                page.insert_text((72, 72), f"Page {page_number}")
            document.save(source)
            document.close()
            storage = StorageService(str(Path(directory) / "app-data"))
            storage.initialize()
            service = DocumentImportService(storage)
            imported = service.import_authorized_paths([str(source)])["documents"][0]["document"]
            first = service.content(imported["id"], 0, 1, page_number=1)
            second = service.content(imported["id"], 0, 1, page_number=2)
            last = service.content(imported["id"], 0, 1, page_number=5)

        self.assertEqual(first["pagination"], {"pageNumber": 1, "pageCount": 5, "hasPrevious": False, "hasNext": True, "previousPageNumber": None, "nextPageNumber": 2})
        self.assertEqual(second["references"][0]["page_number"], 2)
        self.assertTrue(second["pagination"]["hasPrevious"])
        self.assertTrue(second["pagination"]["hasNext"])
        self.assertEqual(last["references"][0]["page_number"], 5)
        self.assertFalse(last["pagination"]["hasNext"])
