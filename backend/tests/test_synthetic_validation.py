"""End-to-end validation tests for the isolated fictional demo scenario.

These tests deliberately inspect ClaimFlow's present boundaries: imported
documents remain evidence/proposals until a separate confirmation workflow
creates domain records.  They must never point at a user's application-data
directory.
"""

from __future__ import annotations

import hashlib
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from group_claims_backend.services.document_import import DocumentImportService
from group_claims_backend.services.matter_review import MatterReviewService
from group_claims_backend.storage.service import StorageService


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
GENERATOR_ROOT = REPOSITORY_ROOT / "scripts" / "validation-sprint"
SOURCE_NAMES = (
    "01_opening_holdings_2025-01-31.xlsx",
    "02_partial_transfer_effective_2025-02-15.pdf",
    "03_claim_right_assignment_effective_2025-02-18.pdf",
    "04_later_holdings_2025-02-28.xlsx",
    "05_contradictory_custodian_statement_2025-02-28.pdf",
    "06_approved_billing_instruction_2025-03-01.pdf",
)


class SyntheticValidationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory(prefix="claimflow-synthetic-validation-")
        self.source_dir = Path(self.temporary.name) / "sources"
        self.source_dir.mkdir()
        subprocess.run([sys.executable, str(GENERATOR_ROOT / "build_synthetic_pdfs.py"), str(self.source_dir)], check=True)
        subprocess.run([sys.executable, str(GENERATOR_ROOT / "build_synthetic_workbooks.py"), str(self.source_dir)], check=True)
        self.storage = StorageService(str(Path(self.temporary.name) / "isolated-app-data"))
        self.assertEqual(self.storage.initialize().status, "ready")
        self.imports = DocumentImportService(self.storage)
        self.response = self.imports.import_authorized_paths(
            [str(self.source_dir / filename) for filename in SOURCE_NAMES]
        )

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_import_vault_hashes_duplicate_and_source_references(self) -> None:
        documents = self.response["documents"]
        self.assertEqual(len(documents), 6)
        self.assertTrue(all(item["ok"] for item in documents))
        self.assertTrue(all(item["document"]["processing_status"] == "parsed" for item in documents))
        duplicate = self.imports.import_authorized_paths([str(self.source_dir / SOURCE_NAMES[0])])
        self.assertTrue(duplicate["documents"][0]["ok"])
        self.assertTrue(duplicate["documents"][0]["duplicate"])

        assert self.storage.database is not None
        assert self.storage.paths is not None
        with self.storage.database.connection() as connection:
            rows = connection.execute(
                "SELECT original_filename, managed_relative_path, content_sha256, imported_at FROM source_documents"
            ).fetchall()
            references = connection.execute(
                "SELECT source_reference FROM extracted_content WHERE source_reference LIKE 'Holdings!%'"
            ).fetchall()
        self.assertEqual(len(rows), 6)
        self.assertTrue(references, "XLSX cells must retain sheet/cell references.")
        for row in rows:
            original = self.source_dir / str(row["original_filename"])
            managed = self.storage.paths.root / str(row["managed_relative_path"])
            self.assertEqual(hashlib.sha256(original.read_bytes()).hexdigest(), str(row["content_sha256"]))
            self.assertEqual(hashlib.sha256(managed.read_bytes()).hexdigest(), str(row["content_sha256"]))
            self.assertNotEqual(str(row["imported_at"])[:10], "2025-01-31")

    def test_pdf_text_preserves_effective_dates_transfer_assignment_and_conflict(self) -> None:
        assert self.storage.database is not None
        with self.storage.database.connection() as connection:
            rows = connection.execute(
                "SELECT value_text FROM extracted_content WHERE content_kind = 'pdf_page'"
            ).fetchall()
        text = "\n".join(str(row["value_text"]) for row in rows).casefold()
        self.assertIn("effective date: 15 february 2025", text)
        self.assertIn("does not assign, transfer", text)
        self.assertIn("claim-rights assignment", text)
        self.assertIn("conflicts with the later holdings statement", text)

    def test_generator_creates_independent_ground_truth_with_expected_events_and_billing(self) -> None:
        workbook = load_workbook(self.source_dir / "ClaimFlow_Synthetic_Ground_Truth.xlsx", data_only=False)
        self.assertEqual(workbook.sheetnames, [
            "Summary", "Documents", "Expected Holdings", "Expected Events",
            "Expected Discrepancies", "Billing Ground Truth", "Actual vs Expected",
        ])
        events = workbook["Expected Events"]
        self.assertEqual(events["A4"].value.date().isoformat(), "2025-02-15")
        self.assertEqual(events["B4"].value, "Partial securities transfer")
        self.assertEqual(events["A5"].value.date().isoformat(), "2025-02-18")
        self.assertEqual(events["B5"].value, "Claim-right assignment")
        billing = workbook["Billing Ground Truth"]
        self.assertEqual(billing["B2"].value, 120000)
        self.assertEqual(billing["E4"].value, "=ROUND(C4*$B$2,2)")
        self.assertEqual(billing["E5"].value, "=ROUND(C5*$B$2,2)")

    def test_document_backed_pro_rata_reconciles_exactly(self) -> None:
        billing = next(item["document"] for item in self.response["documents"] if item["document"]["original_filename"].startswith("06_"))
        service = MatterReviewService(self.storage)
        result = service.ask_question({
            "matterId": "workspace",
            "scope": {"mode": "selected", "documentIds": [billing["id"]]},
            "question": "Calculate the approved pro-rata billing allocation.",
            "acknowledged": True,
            "consentedDocumentIds": [],
        })
        self.assertEqual(result["status"], "completed")
        self.assertEqual(result["model"], "local-document-calculation")
        self.assertIn("USD 34,285.71", str(result["answer"]))
        self.assertIn("USD 85,714.29", str(result["answer"]))
        self.assertIn("USD 120,000.00", str(result["answer"]))
        self.assertTrue(str(result["citations"][0]["sourceReference"]).startswith("page 1 block "))

    def test_import_path_stages_no_confirmed_records_but_persists_source_conflicts(self) -> None:
        """Import may create an issue, but never confirms a business record."""
        assert self.storage.database is not None
        with self.storage.database.connection() as connection:
            holding_count = connection.execute("SELECT COUNT(*) AS count FROM holding_positions").fetchone()["count"]
            assignment_count = connection.execute("SELECT COUNT(*) AS count FROM claim_right_assignments").fetchone()["count"]
            issue_count = connection.execute("SELECT COUNT(*) AS count FROM issues").fetchone()["count"]
        self.assertEqual(holding_count, 0)
        self.assertEqual(assignment_count, 0)
        self.assertEqual(issue_count, 1)


if __name__ == "__main__":
    unittest.main()
