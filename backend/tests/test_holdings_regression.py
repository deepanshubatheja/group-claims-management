"""Regression: a dated holdings comparison question was routed to legal claim status.

The failed question asked for investor, account, ISIN, currency, quantities on
both dates, differences, and legal-transfer caveats.  Its caveat wording
("assignment of legal claim rights") matched the legal-status keyword rule, so
the answer quoted the spreadsheets' legal disclaimer (Holdings!A14) instead of
comparing quantities.  The two workbooks below reproduce the synthetic test
matter's layout cell for cell and are imported through the real parser; the
expected figures are test expectations, not production constants.  Mock
provider only; no live AI requests.
"""

from __future__ import annotations

import json
import os
import tempfile
import time
import unittest
import uuid
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from group_claims_backend.services.ai import AIService, AnswerResult
from group_claims_backend.services.ai_analysis import AIAnalysisService
from group_claims_backend.services.document_import import DocumentImportService
from group_claims_backend.services.matter_review import MatterReviewService
from group_claims_backend.services.question_intent import (
    INTENT_HOLDINGS_COMPARISON, INTENT_LEGAL_STATUS, INTENT_RELATIONSHIP, primary_request_text,
)
from group_claims_backend.storage.service import StorageService

FAILED_QUESTION = (
    "Using the holdings spreadsheets dated 28 March 2023 and 30 April 2023, show for each subfund the investor, "
    "account, ISIN, currency, the quantity on both dates and the quantity difference. Include legal-transfer "
    "caveats: a holding change must not be treated as an assignment of legal claim rights."
)
CAVEAT = (
    "Purpose: compare dated positions. A security holding movement is a factual observation only; it does not by "
    "itself establish a legal-right assignment, litigation status, or billing entitlement."
)
POSITIONS = {
    # subfund: (account, March quantity, April quantity, April note)
    "Horizon Credit Fund A": ("AC-1001", 1_000_000, 600_000, "Reduced by 400,000 from the prior statement"),
    "Horizon Credit Fund B": ("AC-2001", 600_000, 1_000_000, "Increased by 400,000 from the prior statement"),
    "Meridian Income Fund C": ("AC-3001", 400_000, 400_000, "Holding unchanged; participation assessed separately"),
}
_ENVIRONMENT = ("OPENAI_API_KEY", "GROUP_CLAIMS_AI_PROVIDER")


class _RecordingProvider:
    def __init__(self) -> None:
        self.contexts: list[str] = []

    def answer_question(self, question: str, context: str, allowed_citation_keys=()) -> AnswerResult:
        self.contexts.append(context)
        return AnswerResult(
            "Synthetic cited answer.", False, ({"sourceReference": allowed_citation_keys[0], "evidenceSnippet": "x"},),
            "mock", 1, 1, 2,
        )


def _holdings_workbook(path: Path, as_of: datetime, later: bool) -> None:
    """The synthetic test matter's holdings layout: labels, header row 7, records 8-10, total 12, caveat 14."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Holdings"
    sheet["A1"] = "SYNTHETIC TEST DATA - NOT CLIENT HOLDINGS"
    sheet["A3"], sheet["B3"] = "As of date", as_of
    sheet["A4"], sheet["B4"] = "Investor / legal entity", "Northstar Investors LP"
    sheet["A5"], sheet["B5"] = "Matter", "Synthetic Credit Suisse Test Matter"
    for column, heading in zip("ABCDEFGH", (
        "As of date", "Subfund", "Account", "ISIN", "Security", "Currency", "Eligible nominal amount", "Position note",
    )):
        sheet[f"{column}7"] = heading
    for row, (subfund, (account, march, april, note)) in enumerate(POSITIONS.items(), start=8):
        values = (
            as_of, subfund, account, "XS1076957700", "Credit Suisse Group AG 6.25% Perpetual", "USD",
            april if later else march, note if later else "Opening position",
        )
        for column, value in zip("ABCDEFGH", values):
            sheet[f"{column}{row}"] = value
    sheet["A12"], sheet["G12"] = "Control total eligible nominal amount", "=SUM(G8:G10)"
    sheet["A14"] = CAVEAT
    workbook.save(path)


class _SyntheticMatterCase(unittest.TestCase):
    """Both holdings workbooks imported through the real parser, plus three overlapping PDFs."""

    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self._previous_environment = {key: os.environ.get(key) for key in _ENVIRONMENT}
        os.environ["OPENAI_API_KEY"] = "synthetic-key"
        os.environ["GROUP_CLAIMS_AI_PROVIDER"] = "openai"
        self.storage = StorageService(str(Path(self.temporary.name) / "app"))
        self.storage.initialize()
        self.provider = _RecordingProvider()
        analysis = AIAnalysisService(self.storage, provider_factory=lambda: self.provider)  # type: ignore[arg-type]
        self.service = MatterReviewService(self.storage, analysis)
        AIService(self.storage).set_authorization({"authorized": True, "acknowledged": True})
        sources = Path(self.temporary.name) / "sources"
        sources.mkdir()
        march, april = sources / "01_Holdings_2023-03-28.xlsx", sources / "02_Holdings_2023-04-30.xlsx"
        _holdings_workbook(march, datetime(2023, 3, 28), later=False)
        _holdings_workbook(april, datetime(2023, 4, 30), later=True)
        self.original_bytes = {path: path.read_bytes() for path in (march, april)}
        imported = DocumentImportService(self.storage).import_authorized_paths([str(march), str(april)])
        ids = {str(item["document"]["original_filename"]): str(item["document"]["id"]) for item in imported["documents"]}
        self.march, self.april = ids[march.name], ids[april.name]
        # PDFs whose text overlaps heavily with the question (legal rights,
        # assignment, litigation, holdings) but which must not be used.
        self.assignment = self._pdf("03_Legal_Rights_Assignment.pdf", [
            "Assignment of legal claim rights: 400,000 related legal claim rights are assigned from Horizon Credit Fund A "
            "to Horizon Credit Fund B effective 15 April 2023.",
        ])
        self.litigation = self._pdf("04_Litigation_Status_Notice.pdf", [
            "Meridian Income Fund C withdraws from the litigation effective 20 April 2023; its holding remains 400,000.",
        ])
        self.billing = self._pdf("05_Approved_Billing_Instruction.pdf", [
            "Approved fee pool USD 120,000.00 allocated pro rata on 30 April eligible nominal holdings of each account.",
        ])
        self.sources = sources

    def tearDown(self) -> None:
        for key, value in self._previous_environment.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        self.temporary.cleanup()

    def _pdf(self, filename: str, blocks: list[str]) -> str:
        document_id, extraction_id, page_id = str(uuid.uuid4()), str(uuid.uuid4()), str(uuid.uuid4())
        stamp = "2026-01-01T00:00:00+00:00"
        assert self.storage.database
        with self.storage.database.transaction() as connection:
            connection.execute(
                "INSERT INTO source_documents (id, original_filename, managed_relative_path, content_sha256, file_size_bytes, "
                "imported_at, file_format, processing_status) VALUES (?, ?, ?, ?, 1, ?, 'pdf', 'parsed')",
                (document_id, filename, f"documents/{document_id}", (document_id.replace("-", "") * 2)[:64], stamp),
            )
            connection.execute(
                "INSERT INTO document_extractions (id, source_document_id, parser_version, status, started_at, completed_at) "
                "VALUES (?, ?, 'synthetic', 'parsed', ?, ?)", (extraction_id, document_id, stamp, stamp),
            )
            connection.execute(
                "INSERT INTO pdf_page_extractions (id, extraction_id, page_number, extraction_method, raw_text, cleaned_text, "
                "block_count) VALUES (?, ?, 1, 'native_text', ?, ?, ?)",
                (page_id, extraction_id, "\n".join(blocks), "\n".join(blocks), len(blocks)),
            )
            for index, block in enumerate(blocks):
                connection.execute(
                    "INSERT INTO pdf_text_blocks (id, pdf_page_extraction_id, block_index, text, is_table_candidate) VALUES (?, ?, ?, ?, 0)",
                    (str(uuid.uuid4()), page_id, index, block),
                )
        return document_id

    @staticmethod
    def _request(question: str, *, mode: str = "all", document_ids: list[str] | None = None) -> dict[str, object]:
        scope: dict[str, object] = {"mode": mode}
        if document_ids is not None:
            scope["documentIds"] = document_ids
        return {"matterId": "workspace", "scope": scope, "question": question}

    def _ask(self, request: dict[str, object]) -> dict[str, object]:
        preview = self.service.question_preview(request)
        consented = [item["documentId"] for item in preview["documentScope"]["transmittedDocuments"]]
        started = self.service.ask_question({**request, "acknowledged": True, "consentedDocumentIds": consented})
        for _ in range(50):
            status = self.service.question_status({"questionId": started["questionId"]})
            if status["status"] in {"completed", "failed"}:
                return dict(status)
            time.sleep(0.02)
        self.fail("question did not finish")

    @staticmethod
    def _rows(status: dict[str, object]) -> dict[str, dict[str, object]]:
        return {
            next(item["name"] for item in row["entities"] if item["kind"] == "subfund"): row
            for row in status["holdingsComparison"]["rows"]  # type: ignore[index]
        }

    def _assert_expected_comparison(self, status: dict[str, object]) -> None:
        self.assertEqual(status["status"], "completed")
        self.assertEqual(status["intent"], INTENT_HOLDINGS_COMPARISON)
        self.assertEqual(status["model"], "local-holdings-comparison")
        self.assertEqual(self.provider.contexts, [])
        comparison = status["holdingsComparison"]
        assert isinstance(comparison, dict)
        self.assertEqual(comparison["issues"], [])
        self.assertEqual(
            [(item["filename"], item["date"], item["dateBasis"]) for item in comparison["documents"]],
            [("01_Holdings_2023-03-28.xlsx", "2023-03-28", "file name"), ("02_Holdings_2023-04-30.xlsx", "2023-04-30", "file name")],
        )
        rows = self._rows(status)
        self.assertEqual(list(rows), list(POSITIONS))
        expected_change = {"Horizon Credit Fund A": "-400000", "Horizon Credit Fund B": "+400000", "Meridian Income Fund C": "0"}
        for row_number, (subfund, (account, march, april, _note)) in enumerate(POSITIONS.items(), start=8):
            row = rows[subfund]
            self.assertEqual(row["status"], "compared")
            self.assertEqual((row["account"], row["isin"], row["currency"]), (account, "XS1076957700", "USD"))
            self.assertIn({"kind": "investor", "name": "Northstar Investors LP"}, row["entities"])
            self.assertEqual(row["difference"], expected_change[subfund])
            self.assertEqual(Decimal(str(row["difference"])), Decimal(april) - Decimal(march))
            for side, document_id, quantity in (("earlier", self.march, march), ("later", self.april, april)):
                details = row[side]
                self.assertEqual(details["documentId"], document_id)
                self.assertEqual(Decimal(details["quantity"]), Decimal(quantity))
                self.assertEqual(details["quantityReference"], f"Holdings!G{row_number}")
                self.assertEqual(details["isinReference"], f"Holdings!D{row_number}")
                self.assertEqual(details["accountReference"], f"Holdings!C{row_number}")
                self.assertEqual(details["currencyReference"], f"Holdings!F{row_number}")
                self.assertEqual(details["rowDateReference"], f"Holdings!A{row_number}")
                self.assertEqual(details["rowReference"], f"Holdings!A{row_number}:H{row_number}")
                self.assertEqual(set(details["entityReferences"]), {"Holdings!B4", f"Holdings!B{row_number}"})
                # The complete row is returned with its headings, not isolated cells.
                self.assertEqual(
                    [cell["heading"] for cell in details["cells"]],
                    ["As of date", "Subfund", "Account", "ISIN", "Security", "Currency", "Eligible nominal amount", "Position note"],
                )
        self.assertEqual(comparison["totals"], {
            "earlier": "2000000", "later": "2000000", "difference": "0", "currency": "USD", "coverage": "all_records", "note": None,
        })
        for document in comparison["documents"]:
            (table,) = document["tables"]
            self.assertEqual(table["headerReference"], "Holdings!A7:H7")
            self.assertEqual(table["recordReference"], "Holdings!A8:H10")
            self.assertEqual(table["quantityReference"], "Holdings!G8:G10")
            self.assertEqual((table["quantityHeading"], table["quantityHeadingReference"]), ("Eligible nominal amount", "Holdings!G7"))
            self.assertEqual(table["investor"]["name"], "Northstar Investors LP")
            self.assertEqual((table["investor"]["reference"], table["investor"]["labelReference"]), ("Holdings!B4", "Holdings!A4"))
            self.assertEqual((table["statedDate"]["date"], table["statedDate"]["reference"]), (document["date"], "Holdings!B3"))
            control = table["controlTotal"]
            self.assertEqual((control["reference"], control["formula"]), ("Holdings!G12", "=SUM(G8:G10)"))
            self.assertEqual((control["status"], control["recordsTotal"]), ("formula_covers_records", "2000000"))
            self.assertEqual((control["recordRows"], control["comparisonRows"], control["comparisonTotal"], control["reconciled"]), (3, 3, "2000000", True))
        # Holdings!A14 appears only as a caveat, never in a record row.
        self.assertEqual(comparison["caveats"], [{"text": CAVEAT, "sources": [
            {"documentId": self.march, "filename": "01_Holdings_2023-03-28.xlsx", "reference": "Holdings!A14"},
            {"documentId": self.april, "filename": "02_Holdings_2023-04-30.xlsx", "reference": "Holdings!A14"},
        ]}])
        for row in rows.values():
            for side in ("earlier", "later"):
                self.assertNotIn("Holdings!A14", {cell["reference"] for cell in row[side]["cells"]})
        cited = {(citation["documentId"], citation["sourceReference"]) for citation in status["citations"]}  # type: ignore[union-attr]
        for document_id in (self.march, self.april):
            self.assertTrue({
                (document_id, reference) for reference in (
                    "Holdings!B4", "Holdings!G7", "Holdings!G8", "Holdings!G9", "Holdings!G10", "Holdings!G12", "Holdings!A14",
                )
            } <= cited)
        self.assertEqual({citation["documentId"] for citation in status["citations"]}, {self.march, self.april})  # type: ignore[union-attr]
        # The caveat is cited last, after every amount.
        references = [citation["sourceReference"] for citation in status["citations"]]  # type: ignore[union-attr]
        self.assertEqual(references[-2:], ["Holdings!A14", "Holdings!A14"])


class HoldingsComparisonRegressionTests(_SyntheticMatterCase):
    # -- routing ------------------------------------------------------------

    def test_the_failed_question_routes_to_dated_holdings_comparison(self) -> None:
        self.assertEqual(MatterReviewService._question_intent(FAILED_QUESTION).kind, INTENT_HOLDINGS_COMPARISON)
        preview = self.service.question_preview(self._request(FAILED_QUESTION))
        self.assertEqual(preview["intent"], INTENT_HOLDINGS_COMPARISON)
        self.assertTrue(preview["localOnly"])
        self.assertEqual(preview["documentScope"]["transmittedDocuments"], [])

    def test_legal_caveat_wording_does_not_hijack_the_comparison(self) -> None:
        self.assertNotIn("claim", primary_request_text(FAILED_QUESTION))
        for question in (
            "Compare the March 28 and April 30 holdings spreadsheets. Do not infer legal transfers of claims.",
            "Compare holding quantities between the March 28 and April 30 spreadsheets without implying any legal claim assignment.",
            "Compare holdings between the two spreadsheets dated March 28 and April 30; note that holding changes do not establish legal rights.",
        ):
            self.assertEqual(MatterReviewService._question_intent(question).kind, INTENT_HOLDINGS_COMPARISON, question)
        # Affirmative legal questions still route to legal status.
        for question in (
            "Is there explicit evidence of a legal-right assignment?",
            "Which subfund is no longer participating in the litigation?",
            "Did legal claim rights transfer between March 28 and April 30?",
        ):
            self.assertEqual(MatterReviewService._question_intent(question).kind, INTENT_LEGAL_STATUS, question)

    # -- calculation, citations, caveat ---------------------------------------

    def test_expected_quantities_differences_totals_and_cell_citations(self) -> None:
        status = self._ask(self._request(FAILED_QUESTION))
        self._assert_expected_comparison(status)
        answer = str(status["answer"])
        for line in (
            "Investor: Northstar Investors LP.", "ISIN: XS1076957700.", "Currency: USD.",
            "Horizon Credit Fund A / AC-1001: 1,000,000 → 600,000 (change -400,000)",
            "Horizon Credit Fund B / AC-2001: 600,000 → 1,000,000 (change +400,000)",
            "Meridian Income Fund C / AC-3001: 400,000 → 400,000 (change 0)",
            "Totals: 2,000,000 → 2,000,000 (net change 0).",
            "Holdings!G12 holds the formula =SUM(G8:G10) with no stored result",
        ):
            self.assertIn(line, answer)

    def test_the_answer_never_claims_a_legal_claim_assignment(self) -> None:
        status = self._ask(self._request(FAILED_QUESTION))
        answer = str(status["answer"])
        # The legal wording appears only inside the caveat section, quoted from A14.
        body, _, caveat = answer.partition("Legal caveat:")
        self.assertNotRegex(body.casefold(), r"assign|legal|claim|litigation|entitle")
        self.assertIn("does not by itself show a trade, settlement, transfer, or any change in legal claim rights", caveat)
        self.assertIn(f"(01_Holdings_2023-03-28.xlsx Holdings!A14, 02_Holdings_2023-04-30.xlsx Holdings!A14): “{CAVEAT}”", caveat)
        self.assertNotRegex(answer.casefold(), r"\b(?:were|was|have been|has been) (?:assigned|transferred)\b")
        # The assignment notice is never consulted, even though it names the funds.
        self.assertNotIn(self.assignment, {citation["documentId"] for citation in status["citations"]})  # type: ignore[union-attr]

    def test_an_affirmative_legal_part_is_declined_separately(self) -> None:
        status = self._ask(self._request(
            "Compare the holdings quantities in the spreadsheets dated March 28 and April 30 and say whether the legal "
            "claim rights were assigned."
        ))
        self.assertEqual(status["intent"], INTENT_HOLDINGS_COMPARISON)
        self.assertTrue(status["holdingsComparison"]["legalQuestionNotAnswered"])  # type: ignore[index]
        answer = str(status["answer"])
        self.assertIn("They do not establish whether legal claim rights were assigned, transferred or withdrawn", answer)
        # The limitation is stated in this answer; the user is not told to ask again.
        self.assertNotRegex(answer.casefold(), r"ask (?:it|this|again|separately)|on its own")
        self.assertEqual(self.provider.contexts, [])

    # -- scope ------------------------------------------------------------------

    def test_named_dates_narrow_all_documents_and_exclude_every_pdf(self) -> None:
        status = self._ask(self._request(FAILED_QUESTION))
        scope = status["documentScope"]
        self.assertEqual(scope["mode"], "all")  # type: ignore[index]
        self.assertEqual(scope["documentsInScope"], 5)  # type: ignore[index]
        self.assertEqual(
            [item["documentId"] for item in scope["resolvedDocuments"]], [self.march, self.april],  # type: ignore[index]
        )
        self.assertEqual(scope["documentsWithEvidence"], 2)  # type: ignore[index]
        self.assertEqual(scope["transmittedDocuments"], [])  # type: ignore[index]
        cited = {citation["documentId"] for citation in status["citations"]}  # type: ignore[union-attr]
        self.assertFalse(cited & {self.assignment, self.litigation, self.billing})

    def test_selected_documents_only_scope_gives_the_same_comparison(self) -> None:
        with_pdf = self._ask(self._request(
            FAILED_QUESTION, mode="selected", document_ids=[self.march, self.april, self.assignment],
        ))
        self._assert_expected_comparison(with_pdf)
        self.assertEqual(with_pdf["documentScope"]["mode"], "selected")  # type: ignore[index]
        # Without named dates, the two selected spreadsheets are the comparison.
        unnamed = self._ask(self._request(
            "Compare the holdings quantities on both dates and show the differences, with legal caveats.",
            mode="selected", document_ids=[self.march, self.april],
        ))
        self._assert_expected_comparison(unnamed)

    def test_a_named_document_outside_the_selected_scope_is_never_added(self) -> None:
        status = self._ask(self._request(FAILED_QUESTION, mode="selected", document_ids=[self.march, self.assignment]))
        self.assertEqual(status["model"], "local-document-resolution")
        self.assertIn("02_Holdings_2023-04-30.xlsx", str(status["answer"]))
        self.assertIn("not in the current document scope", str(status["answer"]))
        self.assertIsNone(status["holdingsComparison"])
        self.assertEqual(self.provider.contexts, [])

    def test_uncertain_mapping_is_reported_with_cells_not_invented(self) -> None:
        # A later file whose quantity heading is unrecognised cannot be compared.
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Holdings"
        for column, heading in zip("ABC", ("Subfund", "ISIN", "Amount held (bn)")):
            sheet[f"{column}1"] = heading
        sheet["A2"], sheet["B2"], sheet["C2"] = "Horizon Credit Fund A", "XS1076957700", 0.6
        path = self.sources / "03_Holdings_2023-05-31.xlsx"
        workbook.save(path)
        imported = DocumentImportService(self.storage).import_authorized_paths([str(path)])
        may = str(imported["documents"][0]["document"]["id"])
        status = self._ask(self._request(
            "Compare the holdings quantities between the spreadsheets dated April 30 and May 31.",
        ))
        comparison = status["holdingsComparison"]
        assert isinstance(comparison, dict)
        self.assertIn("03_Holdings_2023-05-31.xlsx has no header row with both an ISIN and a quantity column heading.", comparison["issues"])
        self.assertTrue(all(row["status"] == "only_earlier" and row["difference"] is None for row in comparison["rows"]))
        self.assertIsNone(comparison["totals"]["earlier"])
        self.assertNotIn(may, {citation["documentId"] for citation in status["citations"]})  # type: ignore[union-attr]

    def test_original_workbooks_are_not_modified(self) -> None:
        self._ask(self._request(FAILED_QUESTION))
        for path, original in self.original_bytes.items():
            self.assertEqual(path.read_bytes(), original)


RELATIONSHIP_QUESTION = "Which investor is each subfund associated with in these holdings workbooks? Show the exact cells."
LEGAL_RIGHTS_QUESTION = (
    "Horizon Fund A decreased by 400,000 and Horizon Fund B increased by 400,000. Does that mean the legal claim "
    "rights were transferred from Fund A to Fund B?"
)


class _ContextProvider:
    """Records the provider context and cites the last key it was offered."""

    def __init__(self) -> None:
        self.contexts: list[str] = []

    def answer_question(self, question: str, context: str, allowed_citation_keys=()) -> AnswerResult:
        self.contexts.append(context)
        return AnswerResult(
            "The assignment notice separately records an assignment.", False,
            ({"sourceReference": allowed_citation_keys[-1], "evidenceSnippet": "x"},), "mock", 1, 1, 2,
        )


class RelationshipAndLegalGroundingTests(_SyntheticMatterCase):
    """Section-level investor labels and legal-rights answers grounded in dated holding cells."""

    def _seed_summaries_citing_a1(self) -> None:
        # Earlier AI document reviews whose only stored citation is the sheet title.
        assert self.storage.database
        with self.storage.database.transaction() as connection:
            for document_id in (self.march, self.april):
                extraction_id = connection.execute(
                    "SELECT id FROM document_extractions WHERE source_document_id = ?", (document_id,),
                ).fetchone()[0]
                connection.execute(
                    "INSERT INTO ai_processing_jobs (id, source_document_id, extraction_id, provider, model, prompt_version, "
                    "schema_version, status, summary_text, summary_unknown, summary_citations_json, created_at, completed_at) "
                    "VALUES (?, ?, ?, 'openai', 'test', 'v1', 'v1', 'completed', ?, 0, ?, ?, ?)",
                    (str(uuid.uuid4()), document_id, extraction_id,
                     "Synthetic holdings statement listing Horizon Credit Fund A, Horizon Credit Fund B and Meridian Income Fund C.",
                     json.dumps([{"sourceReference": "Holdings!A1", "evidenceSnippet": "SYNTHETIC TEST DATA"}]),
                     "2026-01-01T00:00:00+00:00", "2026-01-01T00:00:00+00:00"),
                )

    # -- A: relationships -------------------------------------------------

    def test_sheet_investor_label_is_associated_with_the_subfund_rows_below_it(self) -> None:
        status = self._ask(self._request(RELATIONSHIP_QUESTION, mode="selected", document_ids=[self.march, self.april]))
        self.assertEqual(status["intent"], INTENT_RELATIONSHIP)
        self.assertEqual(self.provider.contexts, [])
        relationships = status["inventory"]["relationships"]  # type: ignore[index]
        found = {(item["documentId"], item["subject"], item["subjectReference"], item["related"], item["relatedReference"], item["kind"]) for item in relationships}
        expected = {
            (document_id, subfund, f"Holdings!B{row}", "Northstar Investors LP", "Holdings!B4", "document_section")
            for document_id in (self.march, self.april)
            for row, subfund in zip((8, 9, 10), POSITIONS)
        }
        self.assertEqual(found, expected)
        for item in relationships:
            self.assertIn("“Investor / legal entity” (Holdings!A4)", item["basis"])
        investor = next(entity for entity in status["inventory"]["entities"] if entity["name"] == "Northstar Investors LP")  # type: ignore[index]
        self.assertEqual(investor["roles"], ["investor"])
        self.assertEqual({item["sourceReference"] for item in investor["evidence"]}, {"Holdings!B4"})

    def test_association_is_classified_without_asserting_ownership(self) -> None:
        status = self._ask(self._request(RELATIONSHIP_QUESTION, mode="selected", document_ids=[self.march, self.april]))
        answer = str(status["answer"])
        self.assertIn("6 on an investor label above the table that lists the subfunds", answer)
        self.assertIn("None of them by itself proves ownership, separate legal personality", answer)
        self.assertNotRegex(answer.casefold(), r"\b(?:owns|is owned by|is the owner|legal owner)\b")
        self.assertNotIn("documentSection", answer)

    def test_selected_scope_is_preserved_for_relationships(self) -> None:
        status = self._ask(self._request(RELATIONSHIP_QUESTION, mode="selected", document_ids=[self.march]))
        self.assertEqual(status["documentScope"]["mode"], "selected")  # type: ignore[index]
        self.assertEqual(status["documentScope"]["documentsInScope"], 1)  # type: ignore[index]
        self.assertEqual({item["documentId"] for item in status["inventory"]["relationships"]}, {self.march})  # type: ignore[index]
        self.assertEqual({citation["documentId"] for citation in status["citations"]}, {self.march})  # type: ignore[union-attr]

    def test_a_label_applies_only_to_the_table_directly_below_and_header_rows_are_not_pairs(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Register"
        sheet["A1"], sheet["B1"] = "Investor", "Alpha Pension Fund"
        sheet["A3"], sheet["B3"], sheet["C3"] = "Subfund", "ISIN", "Quantity"
        sheet["A4"], sheet["B4"], sheet["C4"] = "Alpha Sub I", "XS0000000001", 10
        # A second table with no investor label of its own.
        sheet["A7"], sheet["B7"], sheet["C7"] = "Subfund", "ISIN", "Quantity"
        sheet["A8"], sheet["B8"], sheet["C8"] = "Beta Sub II", "XS0000000002", 20
        # A two-column header row is a heading row, not "Investor: Quantity".
        sheet["E10"], sheet["F10"] = "Investor", "Quantity"
        path = self.sources / "Register.xlsx"
        workbook.save(path)
        imported = DocumentImportService(self.storage).import_authorized_paths([str(path)])
        register = str(imported["documents"][0]["document"]["id"])
        status = self._ask(self._request(RELATIONSHIP_QUESTION, mode="selected", document_ids=[register]))
        pairs = {(item["subject"], item["related"], item["kind"]) for item in status["inventory"]["relationships"]}  # type: ignore[index]
        self.assertEqual(pairs, {("Alpha Sub I", "Alpha Pension Fund", "document_section")})
        names = {entity["name"] for entity in status["inventory"]["entities"]}  # type: ignore[index]
        self.assertNotIn("Quantity", names)

    # -- B: legal-rights grounding ---------------------------------------

    def test_legal_rights_question_is_supported_by_dated_quantity_cells_and_the_caveat(self) -> None:
        self._seed_summaries_citing_a1()
        status = self._ask(self._request(LEGAL_RIGHTS_QUESTION, mode="selected", document_ids=[self.march, self.april]))
        self.assertEqual((status["intent"], status["model"]), (INTENT_LEGAL_STATUS, "local-legal-status-holdings"))
        self.assertEqual(self.provider.contexts, [])
        cited = {(citation["documentId"], citation["sourceReference"]) for citation in status["citations"]}  # type: ignore[union-attr]
        for document_id in (self.march, self.april):
            self.assertTrue({(document_id, reference) for reference in ("Holdings!G8", "Holdings!G9", "Holdings!B8", "Holdings!B9", "Holdings!A14")} <= cited)
            self.assertNotIn((document_id, "Holdings!A1"), cited)
        # Meridian did not change, so it is not cited as evidence of the movement.
        self.assertNotIn((self.march, "Holdings!G10"), cited)
        # The caveat is supporting context, cited after the factual cells.
        references = [citation["sourceReference"] for citation in status["citations"]]  # type: ignore[union-attr]
        self.assertEqual(references[-2:], ["Holdings!A14", "Holdings!A14"])
        answer = str(status["answer"])
        self.assertIn("Horizon Credit Fund A / AC-1001: 1,000,000 → 600,000 (change -400,000)", answer)
        self.assertIn("Horizon Credit Fund B / AC-2001: 600,000 → 1,000,000 (change +400,000)", answer)
        self.assertIn("does not by itself prove a securities transfer between these accounts or an assignment of legal claim rights", answer)
        self.assertIn("“Horizon Fund A” does not appear under an entity heading", answer)
        self.assertNotIn("“Fund A” does not appear", answer)
        self.assertNotRegex(answer.casefold(), r"\b(?:were|was|have been|has been) (?:assigned|transferred)\b")
        # "Fund A" / "Fund B" resolve to the only names ending that way; Meridian was not asked about.
        rows = self._rows(status)
        self.assertEqual(
            {name: row["difference"] for name, row in rows.items()},
            {"Horizon Credit Fund A": "-400000", "Horizon Credit Fund B": "+400000"},
        )
        self.assertEqual(status["documentScope"]["mode"], "selected")  # type: ignore[index]

    def test_irrelevant_a1_evidence_is_never_the_sole_citation(self) -> None:
        self._seed_summaries_citing_a1()
        self.provider = _ContextProvider()
        self.service._analysis._provider_factory = lambda: self.provider  # type: ignore[attr-defined]
        for question in (LEGAL_RIGHTS_QUESTION, "Were the legal claim rights of Horizon Credit Fund A assigned?"):
            status = self._ask(self._request(question, mode="selected", document_ids=[self.march, self.april]))
            references = {citation["sourceReference"] for citation in status["citations"]}  # type: ignore[union-attr]
            self.assertNotEqual(references, {"Holdings!A1"}, question)
            self.assertNotIn("Holdings!A1", references, question)
        for context in self.provider.contexts:
            self.assertNotIn("SOURCE: Holdings!A1", context)

    def test_with_an_assignment_notice_in_scope_holdings_evidence_is_pinned_first(self) -> None:
        self._seed_summaries_citing_a1()
        self.provider = _ContextProvider()
        self.service._analysis._provider_factory = lambda: self.provider  # type: ignore[attr-defined]
        status = self._ask(self._request(LEGAL_RIGHTS_QUESTION))
        self.assertEqual(status["status"], "completed")
        (context,) = self.provider.contexts
        evidence = context.split("EVIDENCE-LINKED PROVISIONAL FINDINGS:", 1)[1]
        first, *rest = [line for line in evidence.splitlines() if line.startswith("[E")]
        self.assertIn("SOURCE: Holdings!G8; DOCUMENT-BACKED HOLDINGS RECORD", first)
        self.assertIn("1,000,000 → 600,000 (change -400,000)", first)
        self.assertTrue(any("03_Legal_Rights_Assignment.pdf" in line for line in rest))
        self.assertNotIn("SOURCE: Holdings!A1", context)
        cited = {(citation["documentId"], citation["sourceReference"]) for citation in status["citations"]}  # type: ignore[union-attr]
        self.assertTrue({(self.march, "Holdings!G8"), (self.april, "Holdings!G8"), (self.april, "Holdings!A14")} <= cited)
        self.assertIsNotNone(status["holdingsComparison"])
        self.assertTrue(status["holdingsComparison"]["legalQuestionNotAnswered"])  # type: ignore[index]

    def test_verified_holdings_comparison_is_unchanged(self) -> None:
        self._seed_summaries_citing_a1()
        self._assert_expected_comparison(self._ask(self._request(FAILED_QUESTION)))


MULTI_ENTITY_QUESTION = (
    "Using the holdings spreadsheets dated March 28 and April 30, list Horizon Credit Fund A, Horizon Credit Fund B and "
    "Meridian Income Fund C. Compare each subfund's account, ISIN, currency, quantity on both dates and the difference, "
    "including unchanged positions. Explain whether Fund A's decrease and Fund B's increase prove a securities transfer "
    "or a claim-right assignment."
)


class MultiEntityComparisonTests(_SyntheticMatterCase):
    """Every named entity is accounted for, and control totals reconcile only complete comparisons."""

    def _extra_workbooks(self, march_rows: list[tuple[object, ...]], april_rows: list[tuple[object, ...]]) -> list[str]:
        ids = []
        for name, rows in (("07_Holdings_2023-06-30.xlsx", march_rows), ("08_Holdings_2023-07-31.xlsx", april_rows)):
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Holdings"
            for column, heading in zip("ABCDE", ("Subfund", "Account", "ISIN", "Currency", "Quantity")):
                sheet[f"{column}1"] = heading
            for row_number, values in enumerate(rows, start=2):
                for column, value in zip("ABCDE", values):
                    sheet[f"{column}{row_number}"] = value
            path = self.sources / name
            workbook.save(path)
            imported = DocumentImportService(self.storage).import_authorized_paths([str(path)])
            ids.append(str(imported["documents"][0]["document"]["id"]))
        return ids

    def test_the_failed_multi_part_question_compares_all_three_named_funds(self) -> None:
        status = self._ask(self._request(MULTI_ENTITY_QUESTION))
        self._assert_expected_comparison(status)
        comparison = status["holdingsComparison"]
        assert isinstance(comparison, dict)
        self.assertEqual(
            [(entry["text"], entry["name"], entry["match"], entry["positions"]) for entry in comparison["requestedEntities"]],
            [
                ("Horizon Credit Fund A", "Horizon Credit Fund A", "exact", 1),
                ("Horizon Credit Fund B", "Horizon Credit Fund B", "exact", 1),
                ("Meridian Income Fund C", "Meridian Income Fund C", "exact", 1),
                ("Fund A", "Horizon Credit Fund A", "alias", 1),
                ("Fund B", "Horizon Credit Fund B", "alias", 1),
            ],
        )
        answer = str(status["answer"])
        self.assertIn("3 of 3 position(s) had usable quantities on both dates; 2 differ", answer)
        self.assertIn("Meridian Income Fund C / AC-3001: 400,000 → 400,000 (change 0)", answer)
        self.assertIn("Totals: 2,000,000 → 2,000,000 (net change 0).", answer)
        self.assertNotIn("does not appear", answer)
        self.assertNotIn("Compare”", answer)
        # The legal part is answered as a limitation, not a conclusion.
        self.assertIn("does not by itself show a trade, settlement, transfer, or any change in legal claim rights", answer)
        self.assertIn("They do not establish whether legal claim rights were assigned, transferred or withdrawn", answer)
        affirmative = answer.replace("whether legal claim rights were assigned", "")
        self.assertNotRegex(affirmative.casefold(), r"\b(?:were|was|have been|has been) (?:assigned|transferred)\b")

    def test_sentence_boundary_does_not_join_a_name_to_the_next_word(self) -> None:
        from group_claims_backend.services.question_intent import question_entity_names
        names = [variants[0] for variants in question_entity_names(MULTI_ENTITY_QUESTION)]
        self.assertEqual(names, ["Horizon Credit Fund A", "Horizon Credit Fund B", "Meridian Income Fund C", "Fund A", "Fund B"])
        self.assertEqual([variants[0] for variants in question_entity_names("Ask about St. James Capital.")], ["St James Capital"])

    def test_a_named_subset_has_subset_totals_and_no_false_reconciliation(self) -> None:
        status = self._ask(self._request(
            "Compare the holdings of Horizon Credit Fund A and Horizon Credit Fund B between the March 28 and April 30 spreadsheets.",
        ))
        comparison = status["holdingsComparison"]
        assert isinstance(comparison, dict)
        self.assertEqual([row["account"] for row in comparison["rows"]], ["AC-1001", "AC-2001"])
        self.assertEqual(comparison["totals"]["coverage"], "selected_rows")
        self.assertEqual((comparison["totals"]["earlier"], comparison["totals"]["later"]), ("1600000", "1600000"))
        self.assertIn("not the spreadsheets' full holdings totals", comparison["totals"]["note"])
        for document in comparison["documents"]:
            control = document["tables"][0]["controlTotal"]
            self.assertEqual((control["recordRows"], control["comparisonRows"], control["comparisonTotal"], control["reconciled"]), (3, 2, "1600000", False))
            self.assertIn("does not reconcile the comparison totals", control["note"])
        self.assertIn("Totals of the compared positions only (not the full holdings totals): 1,600,000", str(status["answer"]))

    def test_unmatched_and_ambiguous_entities_are_named_and_mark_the_comparison_incomplete(self) -> None:
        first, second = self._extra_workbooks(
            [("Alpha Fund A", "AC-1", "XS0000000001", "EUR", 10), ("Beta Fund A", "AC-2", "XS0000000001", "EUR", 20)],
            [("Alpha Fund A", "AC-1", "XS0000000001", "EUR", 15), ("Beta Fund A", "AC-2", "XS0000000001", "EUR", 20)],
        )
        status = self._ask(self._request(
            "Compare holdings for Fund A and Zeta Growth Fund Z between the June 30 and July 31 spreadsheets.",
            mode="selected", document_ids=[first, second],
        ))
        comparison = status["holdingsComparison"]
        assert isinstance(comparison, dict)
        entries = {entry["text"]: entry for entry in comparison["requestedEntities"]}
        self.assertEqual((entries["Fund A"]["match"], entries["Fund A"]["candidates"]), ("ambiguous", ["Alpha Fund A", "Beta Fund A"]))
        self.assertEqual(entries["Zeta Growth Fund Z"]["match"], "not_found")
        self.assertIn("“Fund A” could refer to more than one entity (Alpha Fund A, Beta Fund A); it was not resolved.", comparison["issues"])
        self.assertIn("“Zeta Growth Fund Z” does not appear under an entity heading in either spreadsheet.", comparison["issues"])
        self.assertEqual(status["model"], "local-holdings-comparison")
        self.assertTrue(status["scope"]["truncated"])  # type: ignore[index]  # reported as incomplete

    def test_missing_rows_duplicates_and_partial_coverage_are_never_totalled(self) -> None:
        first, second = self._extra_workbooks(
            [("Gamma Fund", "AC-7", "XS0000000007", "EUR", 100), ("Delta Fund", "AC-8", "XS0000000008", "EUR", 50)],
            [("Gamma Fund", "AC-7", "XS0000000007", "EUR", 60), ("Gamma Fund", "AC-7", "XS0000000007", "EUR", 40)],
        )
        status = self._ask(self._request(
            "Compare the holdings for Gamma Fund and Delta Fund between the June 30 and July 31 spreadsheets.",
            mode="selected", document_ids=[first, second],
        ))
        comparison = status["holdingsComparison"]
        assert isinstance(comparison, dict)
        statuses = {row["entities"][0]["name"]: (row["status"], row["difference"]) for row in comparison["rows"]}
        self.assertEqual(statuses, {"Gamma Fund": ("duplicate_rows", None), "Delta Fund": ("only_earlier", None)})
        self.assertIsNone(comparison["totals"]["earlier"])
        self.assertEqual({entry["name"]: entry["positions"] for entry in comparison["requestedEntities"]}, {"Gamma Fund": 1, "Delta Fund": 1})
        self.assertIn("could not be compared", str(status["answer"]))


if __name__ == "__main__":
    unittest.main()
