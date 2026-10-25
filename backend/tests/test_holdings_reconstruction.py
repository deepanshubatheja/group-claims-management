"""Follow-up routing and holdings reconstruction in Matter Review.

Regression: after a correct all-document billing answer, the follow-up
"Explain the discrepancy in Horizon Credit Fund B's holdings as of
28 February 2025. Reconstruct its holdings using the opening statement and
transfer notice, …" was narrowed to one spreadsheet.  The as-of date plus the
word "statement" (from "opening statement") were read as a file-date
reference, the later holdings workbook won the tie-break on "holdings", and
the earlier answer was dropped as context because it cited other documents.

The figures below are synthetic test expectations; the mock provider records
any call, and none is expected for a reconstruction.
"""

from __future__ import annotations

import os
import tempfile
import unittest
import uuid
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from group_claims_backend.services.ai import AIService, AnswerResult
from group_claims_backend.services.ai_analysis import AIAnalysisService
from group_claims_backend.services.document_import import DocumentImportService
from group_claims_backend.services.document_references import CandidateDocument, resolve_document_references
from group_claims_backend.services.matter_review import MatterReviewService
from group_claims_backend.services.question_intent import (
    INTENT_CALCULATION, INTENT_HOLDINGS_RECONSTRUCTION, is_holdings_reconstruction,
)
from group_claims_backend.storage.service import StorageService

BILLING_QUESTION = "How much pro-rata is billed for all the subfunds mentioned in these documents?"
FOLLOW_UP = (
    "Explain the discrepancy in Horizon Credit Fund B's holdings as of 28 February 2025. Reconstruct its holdings using "
    "the opening statement and transfer notice, identify every conflicting source, and explain whether the discrepancy "
    "affects the approved pro-rata billing calculation."
)
INSTRUCTION_BLOCKS = [
    "SYNTHETIC - APPROVED PRO-RATA BILLING INSTRUCTION",
    "Issue date: 1 March 2025\nApproved by: Example Investors LP billing committee\nBilling period\n1 February 2025 to 28 February 2025",
    "Fee pool\nUSD 120,000.00",
    "Allocation rule\neligible nominal amount / total eligible nominal amount",
    "Eligible participant 1\nHorizon Credit Fund A (SF-A): 400,000",
    "Eligible participant 2\nHorizon Credit Fund B (SF-B): 1,000,000",
    "Meridian Income Fund C (SF-C) is excluded because it is not participating in the relevant litigation for\nthis billing period.",
]
TRANSFER_BLOCKS = [
    "SYNTHETIC - PARTIAL SECURITIES TRANSFER NOTICE",
    "Issue date: 20 February 2025\nEffective date: 15 February 2025",
    "Example Custody Bank confirms a partial securities transfer of 200,000 nominal\nunits of ISIN XS0123456789, from Horizon\n"
    "Credit Fund A (SF-A), account CUST-01, to Horizon Credit Fund B (SF-B),\naccount CUST-02.",
    "This is a movement of securities only. This notice does not assign, transfer,\nor determine ownership of any legal claim rights.",
]
ASSIGNMENT_BLOCKS = [
    "SYNTHETIC - CLAIM-RIGHTS ASSIGNMENT",
    "Issue date: 21 February 2025\nEffective date: 18 February 2025",
    "Example Investors LP assigns 100 percent of its legal claim rights associated\nwith ISIN XS0123456789 to Northstar Claims SPV.",
    "This assignment concerns legal claim rights only. It does not move securities,\nand does not change the dated holdings of "
    "Horizon Credit Fund A, Horizon Credit Fund B, or Meridian Income\nFund C.",
]
CUSTODIAN_BLOCKS = [
    "SYNTHETIC - CUSTODIAN POSITION STATEMENT",
    "Issue date: 10 March 2025\nStatement as-of date: 28 February 2025",
    "For account CUST-02, Horizon Credit Fund B (SF-B), the recorded holding\nof ISIN XS0123456789 is 900,000 nominal units as of 28 February 2025.",
    "This document conflicts with the later holdings statement stating 1,000,000\nnominal units for the same account. The difference is unresolved.",
]
OPENING = [("CUST-01", "Horizon Credit Fund A (SF-A)", 600_000), ("CUST-02", "Horizon Credit Fund B (SF-B)", 800_000),
           ("CUST-03", "Meridian Income Fund C (SF-C)", 500_000)]
LATER = [("CUST-01", "Horizon Credit Fund A (SF-A)", 400_000), ("CUST-02", "Horizon Credit Fund B (SF-B)", 1_000_000),
         ("CUST-03", "Meridian Income Fund C (SF-C)", 500_000)]
_ENVIRONMENT = ("OPENAI_API_KEY", "GROUP_CLAIMS_AI_PROVIDER")


class _RecordingProvider:
    def __init__(self) -> None:
        self.questions: list[str] = []
        self.contexts: list[str] = []

    def answer_question(self, question: str, context: str, allowed_citation_keys=()) -> AnswerResult:
        self.questions.append(question)
        self.contexts.append(context)
        return AnswerResult("Synthetic.", False, ({"sourceReference": allowed_citation_keys[0], "evidenceSnippet": "x"},), "mock", 1, 1, 2)


class HoldingsReconstructionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self._previous_environment = {key: os.environ.get(key) for key in _ENVIRONMENT}
        os.environ["OPENAI_API_KEY"] = "synthetic-key"
        os.environ["GROUP_CLAIMS_AI_PROVIDER"] = "openai"
        self.storage = StorageService(str(Path(self.temporary.name) / "app"))
        self.storage.initialize()
        self.provider = _RecordingProvider()
        analysis = AIAnalysisService(self.storage, provider_factory=lambda: self.provider)  # type: ignore[arg-type]
        self.service = MatterReviewService(self.storage, analysis)
        AIService(self.storage).set_authorization({"authorized": True, "acknowledged": True})
        self.sources = Path(self.temporary.name) / "sources"
        self.sources.mkdir()

    def tearDown(self) -> None:
        for key, value in self._previous_environment.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        self.temporary.cleanup()

    def _workbook(self, filename: str, as_of: datetime, rows: list[tuple[str, str, int]], isin: str = "XS0123456789") -> str:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Holdings"
        sheet["A1"] = "Synthetic holdings statement"
        sheet["A3"], sheet["B3"] = "Statement date", as_of
        for column, heading in zip("ABCDEF", ("Account", "Subfund", "Security ISIN", "Currency", "Nominal amount", "As-of date")):
            sheet[f"{column}7"] = heading
        for row, (account, subfund, quantity) in enumerate(rows, start=8):
            for column, value in zip("ABCDEF", (account, subfund, isin, "USD", quantity, as_of)):
                sheet[f"{column}{row}"] = value
        path = self.sources / filename
        workbook.save(path)
        original = path.read_bytes()
        imported = DocumentImportService(self.storage).import_authorized_paths([str(path)])
        self.assertEqual(path.read_bytes(), original)
        return str(imported["documents"][0]["document"]["id"])

    def _pdf(self, filename: str, blocks: list[str]) -> str:
        document_id, extraction_id, page_id = str(uuid.uuid4()), str(uuid.uuid4()), str(uuid.uuid4())
        stamp = "2026-01-01T00:00:00+00:00"
        assert self.storage.database
        with self.storage.database.transaction() as connection:
            connection.execute(
                "INSERT INTO source_documents (id, original_filename, managed_relative_path, content_sha256, file_size_bytes, "
                "imported_at, file_format, processing_status) VALUES (?, ?, ?, ?, 1, ?, 'pdf', 'parsed')",
                (document_id, filename, f"documents/{document_id}", (document_id.replace("-", "") * 2)[:64], stamp),
            )
            connection.execute(
                "INSERT INTO document_extractions (id, source_document_id, parser_version, status, started_at, completed_at) "
                "VALUES (?, ?, 'synthetic', 'parsed', ?, ?)", (extraction_id, document_id, stamp, stamp),
            )
            connection.execute(
                "INSERT INTO pdf_page_extractions (id, extraction_id, page_number, extraction_method, raw_text, cleaned_text, "
                "block_count) VALUES (?, ?, 1, 'native_text', ?, ?, ?)",
                (page_id, extraction_id, "\n".join(blocks), "\n".join(blocks), len(blocks)),
            )
            for index, block in enumerate(blocks):
                connection.execute(
                    "INSERT INTO pdf_text_blocks (id, pdf_page_extraction_id, block_index, text, is_table_candidate) VALUES (?, ?, ?, ?, 0)",
                    (str(uuid.uuid4()), page_id, index, block),
                )
        return document_id

    def _matter(self, *, transfer: bool = True, instruction: bool = True) -> dict[str, str]:
        ids = {
            "opening": self._workbook("01_opening_holdings_2025-01-31.xlsx", datetime(2025, 1, 31), OPENING),
            "assignment": self._pdf("03_claim_right_assignment_effective_2025-02-18.pdf", ASSIGNMENT_BLOCKS),
            "later": self._workbook("04_later_holdings_2025-02-28.xlsx", datetime(2025, 2, 28), LATER),
            "custodian": self._pdf("05_contradictory_custodian_statement_2025-02-28.pdf", CUSTODIAN_BLOCKS),
        }
        if transfer:
            ids["transfer"] = self._pdf("02_partial_transfer_effective_2025-02-15.pdf", TRANSFER_BLOCKS)
        if instruction:
            ids["instruction"] = self._pdf("06_approved_billing_instruction_2025-03-01.pdf", INSTRUCTION_BLOCKS)
        return ids

    def _ask(self, question: str, *, follow_up: str | None = None, document_ids: list[str] | None = None) -> dict:
        scope: dict[str, object] = {"mode": "all"} if document_ids is None else {"mode": "selected", "documentIds": document_ids}
        request: dict[str, object] = {"matterId": "workspace", "scope": scope, "question": question}
        if follow_up:
            request["followUpQuestionId"] = follow_up
        preview = self.service.question_preview(request)
        consented = [item["documentId"] for item in preview["documentScope"]["transmittedDocuments"]]
        status = self.service.ask_question({**request, "acknowledged": True, "consentedDocumentIds": consented})
        self.assertEqual(status["status"], "completed")
        return status

    def _section(self, answer: str, heading: str) -> str:
        """The lines of one answer section, up to the next blank line."""
        start = answer.index(heading + ":\n") + len(heading) + 2
        end = answer.find("\n\n", start)
        return answer[start:] if end < 0 else answer[start:end]

    def _cited(self, status: dict) -> set[tuple[str, str]]:
        return {(citation["documentId"], citation["sourceReference"]) for citation in status["citations"]}

    # -- the reported follow-up -------------------------------------------

    def test_exact_follow_up_keeps_all_documents_and_reconstructs_with_citations(self) -> None:
        ids = self._matter()
        first = self._ask(BILLING_QUESTION)
        self.assertEqual(first["intent"], INTENT_CALCULATION)
        status = self._ask(FOLLOW_UP, follow_up=first["questionId"])
        scope = status["documentScope"]

        # Routing and scope: all six documents stay in scope; nothing is named or sent.
        self.assertEqual(status["intent"], INTENT_HOLDINGS_RECONSTRUCTION)
        self.assertEqual(status["model"], "local-holdings-reconstruction")
        self.assertEqual(scope["mode"], "all")
        self.assertEqual(scope["documentsInScope"], 6)
        self.assertEqual(scope["resolvedDocuments"], [])
        self.assertEqual(scope["transmittedDocuments"], [])
        self.assertFalse(any("not reused" in note for note in scope["notes"]))
        self.assertEqual(self.provider.questions, [])
        self.assertIsNone(status["billingAllocation"])  # not a repeat of the allocation table

        answer = status["answer"]
        sections = ["Documented facts:", "Derived arithmetic:", "Discrepancies:", "Billing effect (provisional):",
                    "Not applied to the securities balance:", "Limits of this answer:"]
        self.assertEqual([answer.index(heading) for heading in sections], sorted(answer.index(heading) for heading in sections))

        facts = self._section(answer, "Documented facts").splitlines()
        self.assertEqual(len(facts), 4)
        self.assertTrue(facts[0].startswith("- 31 Jan 2025 · holdings record: 800,000 nominal, account CUST-02 (01_opening"))
        self.assertTrue(facts[1].startswith("- 15 Feb 2025 (effective) · securities transfer of 200,000 nominal into Horizon Credit Fund B (SF-B)"))
        self.assertIn("1,000,000 nominal, account CUST-02 (04_later_holdings_2025-02-28.xlsx, Holdings!E9)", facts[2] + facts[3])
        self.assertIn("holding statement: 900,000 nominal (05_contradictory_custodian_statement_2025-02-28.pdf, page 1 block 2)", facts[2] + facts[3])
        self.assertIn(
            "800,000 (opening, 31 Jan 2025) + 200,000 (transfer in, 15 Feb 2025) = 1,000,000 nominal as of 28 Feb 2025.",
            self._section(answer, "Derived arithmetic"),
        )
        discrepancies = self._section(answer, "Discrepancies")
        self.assertIn("(a difference of 100,000)", discrepancies)
        self.assertIn("page 1 block 2 records 900,000 nominal, 100,000 less than the derived 1,000,000", discrepancies)
        self.assertIn("page 1 block 3: the source itself states the conflict", discrepancies)
        self.assertIn("Unresolved", discrepancies)

        billing = self._section(answer, "Billing effect (provisional)")
        self.assertIn("uses 1,000,000 nominal as Horizon Credit Fund B (SF-B)'s eligible amount", billing)
        self.assertIn("USD 85,714.29 (71.4286% of the USD 120,000.00 fee pool)", billing)
        self.assertIn("agrees with 04_later_holdings_2025-02-28.xlsx, Holdings!E9; the derived balance", billing)
        self.assertIn("differs from 05_contradictory_custodian_statement_2025-02-28.pdf, page 1 block 2 (900,000)", billing)
        self.assertIn("should not be invoiced until it is resolved", billing)
        self.assertIn("did not substitute another figure or recalculate", billing)

        # Every conflicting source and every input is cited, from the right documents.
        cited = self._cited(status)
        for expected in [
            (ids["opening"], "Holdings!E9"), (ids["transfer"], "page 1 block 2"), (ids["later"], "Holdings!E9"),
            (ids["custodian"], "page 1 block 2"), (ids["custodian"], "page 1 block 3"), (ids["instruction"], "page 1 block 5"),
            (ids["assignment"], "page 1 block 3"), (ids["transfer"], "page 1 block 3"),
        ]:
            self.assertIn(expected, cited)
        self.assertEqual(scope["documentsWithEvidence"], 6)

    def test_reconstruction_never_sends_and_leaves_the_approved_allocation_unchanged(self) -> None:
        self._matter()
        before = self._ask(BILLING_QUESTION)["billingAllocation"]
        self._ask(FOLLOW_UP, follow_up=self._ask(BILLING_QUESTION)["questionId"])
        after = self._ask(BILLING_QUESTION)["billingAllocation"]
        self.assertEqual(before, after)
        self.assertEqual(self.provider.questions, [])

    def test_follow_up_without_names_resolves_the_earlier_conflict_and_regrounds(self) -> None:
        ids = self._matter()
        first = self._ask(BILLING_QUESTION)
        status = self._ask("Explain that discrepancy and reconstruct the holding.", follow_up=first["questionId"])
        self.assertEqual(status["intent"], INTENT_HOLDINGS_RECONSTRUCTION)
        self.assertTrue(any("“Horizon Credit Fund B”" in note for note in status["documentScope"]["notes"]))
        self.assertIn("Horizon Credit Fund B (SF-B)", status["answer"])
        # Facts come from the documents, not from the earlier answer text.
        self.assertIn((ids["custodian"], "page 1 block 2"), self._cited(status))
        self.assertNotIn("Billing effect", status["answer"])  # not asked this time

    # -- scope: descriptive phrases and fact dates are not document references

    def test_descriptive_phrases_and_as_of_dates_do_not_narrow_scope(self) -> None:
        documents = [
            CandidateDocument("1", "01_opening_holdings_2025-01-31.xlsx", "xlsx", True),
            CandidateDocument("4", "04_later_holdings_2025-02-28.xlsx", "xlsx", True),
            CandidateDocument("5", "05_contradictory_custodian_statement_2025-02-28.pdf", "pdf", True),
        ]
        for question in [
            FOLLOW_UP,
            "What did the custodian statement record as of 28 February 2025?",
            "Using the opening statement, the transfer notice and the billing instruction, what was held effective 15 February 2025?",
            "Identify every conflicting source for the holdings on 28 February 2025.",
            # One dated document named alongside documents it only describes is a mention, not a restriction.
            "Compare the 28 February 2025 custodian statement with the transfer notice.",
        ]:
            with self.subTest(question=question):
                self.assertFalse(resolve_document_references(question, documents).requested)

    def test_explicit_dated_or_named_documents_still_narrow_scope(self) -> None:
        documents = [
            CandidateDocument("1", "01_opening_holdings_2025-01-31.xlsx", "xlsx", True),
            CandidateDocument("4", "04_later_holdings_2025-02-28.xlsx", "xlsx", True),
            CandidateDocument("5", "05_contradictory_custodian_statement_2025-02-28.pdf", "pdf", True),
        ]
        for question, expected in [
            ("What does the 28 February 2025 custodian statement record for Fund B?", ["5"]),
            ("In the spreadsheet dated 28 February 2025, what is recorded for Fund B?", ["4"]),
            ("Compare the 31 January 2025 and 28 February 2025 holdings spreadsheets.", ["1", "4"]),
            ("Using only 01_opening_holdings_2025-01-31.xlsx, what is recorded?", ["1"]),
        ]:
            with self.subTest(question=question):
                constraint = resolve_document_references(question, documents)
                self.assertTrue(constraint.requested)
                self.assertEqual([item.document_id for item in constraint.resolved], expected)
                self.assertEqual(constraint.problems, ())

    def test_legitimate_single_document_question_uses_only_that_document(self) -> None:
        ids = self._matter()
        status = self._ask(
            "Using only 01_opening_holdings_2025-01-31.xlsx, reconstruct Horizon Credit Fund B's holdings as of 28 February 2025."
        )
        self.assertEqual([item["documentId"] for item in status["documentScope"]["resolvedDocuments"]], [ids["opening"]])
        self.assertEqual({citation["documentId"] for citation in status["citations"]}, {ids["opening"]})
        answer = status["answer"]
        self.assertTrue(answer.startswith("No record in scope states Horizon Credit Fund B (SF-B)'s holding in XS0123456789 on 28 Feb 2025."))
        self.assertIn("that is arithmetic, not a documented balance", answer)
        self.assertIn("No record in scope is dated 28 Feb 2025, so nothing could be compared.", answer)

    def test_multi_document_ai_question_is_not_restricted_by_its_date(self) -> None:
        self._matter()
        preview = self.service.question_preview({
            "matterId": "workspace", "scope": {"mode": "all"},
            "question": "What do the opening statement and the transfer notice record for Horizon Credit Fund B as of 15 February 2025?",
        })
        self.assertFalse(preview["localOnly"])  # the provider path, with every document available to retrieval
        self.assertEqual(preview["documentScope"]["resolvedDocuments"], [])
        self.assertEqual(preview["documentScope"]["documentsInScope"], 6)
        self.assertGreater(preview["documentScope"]["documentsWithEvidence"], 1)

    # -- reconstruction rules --------------------------------------------------

    def test_routing_prefers_reconstruction_over_a_billing_keyword(self) -> None:
        self.assertTrue(is_holdings_reconstruction(FOLLOW_UP))
        self.assertTrue(is_holdings_reconstruction("Which sources conflict on Fund B's holding?"))
        self.assertFalse(is_holdings_reconstruction(BILLING_QUESTION))
        self.assertFalse(is_holdings_reconstruction("Compare the holdings; do not treat a conflict as a legal assignment."))
        self.assertEqual(MatterReviewService._question_intent(FOLLOW_UP).kind, INTENT_HOLDINGS_RECONSTRUCTION)
        self.assertEqual(MatterReviewService._question_intent(BILLING_QUESTION).kind, INTENT_CALCULATION)

    def test_chronological_window_applies_only_transfers_between_opening_and_date(self) -> None:
        self._matter()
        self._pdf("07_transfer_effective_2025-03-05.pdf", [
            "Effective date: 5 March 2025",
            "A securities transfer of 50,000 nominal units of ISIN XS0123456789 from Horizon Credit Fund B (SF-B), account CUST-02, "
            "to Example Fund Z.",
        ])
        self._pdf("08_transfer_effective_2025-01-31.pdf", [
            "Effective date: 31 January 2025",
            "A securities transfer of 10,000 nominal units of ISIN XS0123456789 from Example Fund Y to Horizon Credit Fund B (SF-B).",
        ])
        answer = self._ask("Reconstruct Horizon Credit Fund B's holdings as of 28 February 2025.")["answer"]
        self.assertIn("= 1,000,000 nominal as of 28 Feb 2025.", answer)
        limits = self._section(answer, "Limits of this answer")
        self.assertIn("The transfer effective 5 Mar 2025 (07_transfer_effective_2025-03-05.pdf, page 1 block 1) is after 28 Feb 2025", limits)
        self.assertIn("The transfer effective 31 Jan 2025 (08_transfer_effective_2025-01-31.pdf, page 1 block 1) falls on the opening date", limits)
        self.assertNotIn("Billing effect", answer)

    def test_claim_rights_are_listed_but_never_change_the_balance(self) -> None:
        self._matter()
        answer = self._ask("Reconstruct Horizon Credit Fund A's holdings as of 28 February 2025.")["answer"]
        self.assertIn("600,000 (opening, 31 Jan 2025) − 200,000 (transfer out, 15 Feb 2025) = 400,000 nominal", answer)
        self.assertIn("out of Horizon Credit Fund A (SF-A) to Horizon Credit Fund B (SF-B)", answer)
        not_applied = self._section(answer, "Not applied to the securities balance")
        self.assertIn("03_claim_right_assignment_effective_2025-02-18.pdf, page 1 block 3: concerns legal claim rights", not_applied)
        self.assertIn("02_partial_transfer_effective_2025-02-15.pdf, page 1 block 3", not_applied)
        self.assertIn("None found among the records in scope dated 28 Feb 2025.", answer)
        self.assertNotIn("Northstar", self._section(answer, "Documented facts"))

    def test_billing_context_for_agreeing_and_excluded_subjects(self) -> None:
        self._matter()
        fund_a = self._ask("Is there a discrepancy in Horizon Credit Fund A's holdings, and does it affect the pro-rata billing?")["answer"]
        self.assertIn("does not change the approved allocation", fund_a)
        self.assertNotIn("should not be invoiced", fund_a)
        fund_c = self._ask("Is there a discrepancy in Meridian Income Fund C's holdings, and does it affect the pro-rata billing?")["answer"]
        self.assertIn("The approved billing instruction excludes Meridian Income Fund C (SF-C) because it is not participating", fund_c)
        self.assertIn("its allocation is 0.00 under that instruction", fund_c)

    def test_missing_transfer_evidence_is_reported_not_assumed(self) -> None:
        ids = self._matter(transfer=False, instruction=False)
        status = self._ask(FOLLOW_UP)
        answer = status["answer"]
        self.assertIn("800,000 (opening, 31 Jan 2025) = 800,000 nominal as of 28 Feb 2025.", answer)
        self.assertIn("no other transfer for Horizon Credit Fund B (SF-B) between 31 Jan 2025 and 28 Feb 2025 is documented", answer)
        self.assertIn("200,000 more than the derived 800,000", answer)
        self.assertIn("No approved billing instruction in scope", answer)
        self.assertNotIn("transfer of", self._section(answer, "Documented facts"))
        # Only documents actually in scope are cited.
        self.assertTrue({citation["documentId"] for citation in status["citations"]} <= set(ids.values()))

    def test_several_securities_need_an_isin(self) -> None:
        self._workbook("01_opening_holdings_2025-01-31.xlsx", datetime(2025, 1, 31), OPENING)
        self._workbook("04_later_holdings_2025-02-28.xlsx", datetime(2025, 2, 28), LATER, isin="XS9999999995")
        answer = self._ask("Reconstruct Horizon Credit Fund B's holdings as of 28 February 2025.")["answer"]
        self.assertIn("cover 2 securities (XS0123456789, XS9999999995); name the ISIN", answer)
        named = self._ask("Reconstruct Horizon Credit Fund B's holdings in XS9999999995 as of 28 February 2025.")["answer"]
        self.assertIn("No holding of Horizon Credit Fund B (SF-B) is documented before 28 Feb 2025", named)

    def test_question_without_records_falls_back_to_the_existing_route(self) -> None:
        self._matter()
        status = self._ask("Explain the discrepancy in Omega Fund Z's holdings.")
        self.assertNotEqual(status["intent"], INTENT_HOLDINGS_RECONSTRUCTION)


if __name__ == "__main__":
    unittest.main()
