"""Deterministic, local document import and extraction services."""

from __future__ import annotations

import csv
import json
import os
import re
import sqlite3
import uuid
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from ..storage.documents import DocumentMetadataRepository, DocumentRecord
from ..storage.service import StorageService
from ..storage.vault import VaultImportError, VaultStorageError

PARSER_VERSION = "deterministic-v4"
SUPPORTED_FORMATS = {".xlsx": "xlsx", ".xls": "xls", ".csv": "csv", ".pdf": "pdf"}
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_IMPORT_FILES = 20
MAX_EXTRACTED_ITEMS = 200_000
MAX_VALUE_CHARS = 100_000
MAX_PREVIEW_VALUE_CHARS = 8_192
MAX_PREVIEW_RESPONSE_BYTES = 48_000
# The protocol response cap is 60 KiB. Keep a deliberately smaller grid window
# and bound presentation values so a dense worksheet cannot make a preview fail.
MAX_GRID_ROWS = 10
MAX_GRID_COLUMNS = 10
MAX_GRID_VALUE_CHARS = 64
MAX_GRID_RESPONSE_BYTES = 30_000
MAX_PDF_PAGES = 100
MAX_OCR_PIXELS = 20_000_000
OCR_RENDER_SCALE = 2
OCR_PAGE_TIMEOUT_SECONDS = 30


class DocumentImportError(RuntimeError):
    """A controlled document import failure suitable for the desktop UI."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


class ExtractionUnsupported(DocumentImportError):
    pass


@dataclass(frozen=True)
class ExtractedItem:
    content_kind: str
    source_reference: str
    value_text: str | None
    value_type: str | None = None
    formula_text: str | None = None
    sheet_name: str | None = None
    row_number: int | None = None
    column_number: int | None = None
    page_number: int | None = None
    worksheet_index: int | None = None
    is_empty: bool = False


@dataclass(frozen=True)
class PdfTextBlock:
    block_index: int
    text: str
    x0: float | None = None
    y0: float | None = None
    x1: float | None = None
    y1: float | None = None
    confidence: float | None = None
    is_table_candidate: bool = False


@dataclass(frozen=True)
class PdfPageExtraction:
    page_number: int
    extraction_method: str
    raw_text: str | None
    cleaned_text: str | None
    quality_warning: str | None
    ocr_language: str | None
    ocr_engine_version: str | None
    mean_confidence: float | None
    blocks: tuple[PdfTextBlock, ...]


@dataclass(frozen=True)
class ParseResult:
    items: tuple[ExtractedItem, ...]
    pdf_pages: tuple[PdfPageExtraction, ...] = ()


def _now() -> str:
    return datetime.now(UTC).isoformat()


def _short_value(value: object) -> str | None:
    if value is None:
        return None
    text = str(value)
    if len(text) > MAX_VALUE_CHARS:
        raise DocumentImportError("CONTENT_TOO_LARGE", "A document value exceeds the safe preview limit.")
    return text


class DeterministicDocumentParser:
    """Parses supported files without interpreting legal or financial meaning."""

    def parse(self, path: Path, file_format: str) -> ParseResult:
        if file_format == "csv":
            return ParseResult(tuple(self._parse_csv(path)))
        if file_format == "xlsx":
            return ParseResult(tuple(self._parse_xlsx(path)))
        if file_format == "xls":
            return ParseResult(tuple(self._parse_xls(path)))
        if file_format == "pdf":
            return self._parse_pdf(path)
        raise ExtractionUnsupported("UNSUPPORTED_FORMAT", "This document format is not supported.")

    @staticmethod
    def _require(package: str) -> None:
        raise DocumentImportError(
            "PARSER_UNAVAILABLE",
            f"The local {package} parser is not installed. Install the documented backend dependencies.",
        )

    def _parse_csv(self, path: Path) -> list[ExtractedItem]:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("utf-8")
            except UnicodeDecodeError as error:
                raise DocumentImportError("INVALID_ENCODING", "CSV files must use UTF-8 encoding.") from error
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(text.splitlines(), dialect))
        items: list[ExtractedItem] = []
        for row_number, row in enumerate(rows, start=1):
            for column_number, value in enumerate(row, start=1):
                items.append(
                    ExtractedItem(
                        "csv_cell", f"CSV!{self._column_letters(column_number)}{row_number}", value,
                        "text", sheet_name="CSV", row_number=row_number, column_number=column_number,
                        worksheet_index=1, is_empty=value == "",
                    )
                )
        self._check_item_limit(items)
        return items

    def _parse_xlsx(self, path: Path) -> list[ExtractedItem]:
        try:
            from openpyxl import load_workbook
        except ImportError:
            self._require("openpyxl")
        items: list[ExtractedItem] = []
        try:
            # Managed vault files intentionally have generated IDs without source extensions.
            # Passing a stream avoids openpyxl's filename-extension gate while still parsing
            # the immutable copied bytes. Keep the stream open for read-only iteration.
            with path.open("rb") as source:
                workbook = load_workbook(source, read_only=True, data_only=False)
                try:
                    for worksheet_index, worksheet in enumerate(workbook.worksheets, start=1):
                        # Some valid workbooks omit the optional worksheet
                        # dimension metadata.  ``openpyxl`` then exposes both
                        # dimensions as ``None`` in read-only mode even though
                        # the worksheet contains cells.  Ask it to determine
                        # the dimensions by scanning the sheet before applying
                        # our extraction limit.  This keeps the memory-safe
                        # streaming parser while accepting standards-compliant
                        # workbooks produced by other spreadsheet tools.
                        if worksheet.max_row is None or worksheet.max_column is None:
                            worksheet.reset_dimensions()
                            worksheet.calculate_dimension(force=True)
                        cells = (worksheet.max_row or 0) * (worksheet.max_column or 0)
                        if cells + len(items) > MAX_EXTRACTED_ITEMS:
                            raise DocumentImportError("CONTENT_TOO_LARGE", "The workbook exceeds the safe extraction limit.")
                        for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                            for column_number, cell in enumerate(row, start=1):
                                # In read-only workbooks, blank positions are EmptyCell objects
                                # without coordinate/row/column fields. Enumeration preserves
                                # their original grid position without fabricating a value.
                                coordinate = f"{self._column_letters(column_number)}{row_number}"
                                value = getattr(cell, "value", None)
                                cell_type = getattr(cell, "data_type", "empty")
                                formula = _short_value(value) if cell_type == "f" else None
                                items.append(ExtractedItem(
                                    "spreadsheet_cell", f"{worksheet.title}!{coordinate}", _short_value(value),
                                    cell_type, formula, worksheet.title, row_number, column_number, None, worksheet_index,
                                    is_empty=value is None,
                                ))
                finally:
                    workbook.close()
        except DocumentImportError:
            raise
        except Exception as error:
            raise DocumentImportError("INVALID_XLSX", "The Excel workbook could not be read.") from error
        return items

    def _parse_xls(self, path: Path) -> list[ExtractedItem]:
        try:
            import xlrd
        except ImportError:
            self._require("xlrd")
        try:
            workbook = xlrd.open_workbook(path, on_demand=True)
        except Exception as error:
            raise DocumentImportError("INVALID_XLS", "The legacy Excel workbook could not be read.") from error
        items: list[ExtractedItem] = []
        try:
            for sheet_index in range(workbook.nsheets):
                sheet = workbook.sheet_by_index(sheet_index)
                if sheet.nrows * sheet.ncols + len(items) > MAX_EXTRACTED_ITEMS:
                    raise DocumentImportError("CONTENT_TOO_LARGE", "The workbook exceeds the safe extraction limit.")
                for row_index in range(sheet.nrows):
                    for column_index in range(sheet.ncols):
                        cell = sheet.cell(row_index, column_index)
                        row_number, column_number = row_index + 1, column_index + 1
                        items.append(ExtractedItem(
                            "spreadsheet_cell", f"{sheet.name}!R{row_number}C{column_number}",
                            _short_value(cell.value), str(cell.ctype), sheet_name=sheet.name,
                            row_number=row_number, column_number=column_number, worksheet_index=sheet_index + 1, is_empty=cell.value == "",
                        ))
        finally:
            workbook.release_resources()
        return items

    def _parse_pdf(self, path: Path) -> ParseResult:
        """Extract PDF text page-by-page, using local OCR only where native text is absent.

        PyMuPDF provides ordered blocks with coordinates.  We preserve those raw
        blocks and deliberately avoid guessing a semantic table layout.  OCR is
        a page-level fallback, never a replacement for usable native text.
        """
        try:
            import pymupdf
        except ImportError:
            self._require("PyMuPDF")
        try:
            document = pymupdf.open(path)
        except Exception as error:
            raise DocumentImportError("INVALID_PDF", "The PDF could not be read.") from error
        try:
            if document.needs_pass:
                raise ExtractionUnsupported("PDF_ENCRYPTED", "The PDF is encrypted and cannot be parsed locally.")
            if not document.page_count:
                raise DocumentImportError("INVALID_PDF", "The PDF contains no pages.")
            if document.page_count > MAX_PDF_PAGES:
                raise DocumentImportError("CONTENT_TOO_LARGE", "The PDF exceeds the safe page-count limit.")

            items: list[ExtractedItem] = []
            pages: list[PdfPageExtraction] = []
            for page_number in range(1, document.page_count + 1):
                page = document.load_page(page_number - 1)
                native_blocks = self._native_pdf_blocks(page)
                if native_blocks:
                    raw_text = "\n".join(block.text for block in native_blocks)
                    cleaned_text = self._clean_pdf_text(raw_text)
                    pages.append(PdfPageExtraction(
                        page_number, "native_text", raw_text, cleaned_text,
                        self._pdf_quality_warning(native_blocks), None, None, None, tuple(native_blocks),
                    ))
                    items.append(ExtractedItem("pdf_page", f"page {page_number}", _short_value(cleaned_text),
                        "native_text", page_number=page_number, is_empty=not cleaned_text.strip()))
                    continue

                ocr_page = self._ocr_pdf_page(page, page_number)
                pages.append(ocr_page)
                items.append(ExtractedItem("pdf_page", f"page {page_number}", _short_value(ocr_page.cleaned_text),
                    ocr_page.extraction_method, page_number=page_number,
                    is_empty=not bool((ocr_page.cleaned_text or "").strip())))
        except ExtractionUnsupported:
            raise
        except DocumentImportError:
            raise
        except Exception as error:
            raise DocumentImportError("INVALID_PDF", "PDF text could not be extracted safely.") from error
        finally:
            document.close()

        if not items or all(item.is_empty for item in items):
            if any(page.extraction_method == "unavailable" for page in pages):
                raise ExtractionUnsupported("OCR_UNAVAILABLE", "OCR is unavailable for PDF pages without extractable text.")
            raise ExtractionUnsupported("PDF_OCR_REQUIRED", "No usable text was extracted from this PDF.")
        self._check_item_limit(items)
        return ParseResult(tuple(items), tuple(pages))

    @staticmethod
    def _clean_pdf_text(text: str) -> str:
        """Normalise presentation whitespace while retaining the raw text separately."""
        # A hyphen immediately followed by a lower-case continuation is a common
        # PDF line-wrap convention.  The unmodified source block remains stored.
        unwrapped = re.sub(r"(?<=[A-Za-zÄÖÜäöüß])-\s*\n(?=[a-zäöüß])", "", text)
        return "\n".join(" ".join(line.split()) for line in unwrapped.splitlines() if line.strip())

    @staticmethod
    def _native_pdf_blocks(page: object) -> list[PdfTextBlock]:
        # ``sort=True`` is a coordinate-based, deterministic reading order. It is
        # still only a presentation aid; the original PDF remains authoritative.
        blocks: list[PdfTextBlock] = []
        for index, block in enumerate(page.get_text("blocks", sort=True)):  # type: ignore[attr-defined]
            x0, y0, x1, y1, text, _block_number, block_type = block[:7]
            if block_type != 0 or not isinstance(text, str) or not text.strip():
                continue
            blocks.append(PdfTextBlock(index, text, float(x0), float(y0), float(x1), float(y1)))
        return blocks

    @staticmethod
    def _pdf_quality_warning(blocks: Sequence[PdfTextBlock]) -> str | None:
        # A visual table can only be treated as a table when its structure is
        # reliable. Keep blocks rather than inventing rows/columns otherwise.
        warnings: list[str] = []
        if len(blocks) >= 3 and len({round(block.x0 or 0, 1) for block in blocks}) >= 2:
            warnings.append("Multiple text columns or a table may affect reading order; inspect source blocks.")
        if any(re.search(r"(?<=[A-Za-zÄÖÜäöüß])-\s*\n(?=[a-zäöüß])", block.text) for block in blocks):
            warnings.append("Cleaned text joins likely end-of-line hyphenation; raw blocks remain available.")
        return " ".join(warnings) or None

    def _ocr_pdf_page(self, page: object, page_number: int) -> PdfPageExtraction:
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            return PdfPageExtraction(page_number, "unavailable", None, None,
                "OCR dependencies are not installed for this backend.", None, None, None, ())
        try:
            # A packaged ClaimFlow backend receives these paths only from the
            # trusted Electron main process. Development keeps the normal local
            # Tesseract discovery behaviour, so no user-installed path is
            # required by the distributed application.
            configured_command = os.environ.get("GROUP_CLAIMS_TESSERACT_CMD")
            if configured_command:
                command_path = Path(configured_command)
                if not command_path.is_file():
                    return PdfPageExtraction(page_number, "unavailable", None, None,
                        "The bundled local OCR runtime is unavailable.", None, None, None, ())
                pytesseract.pytesseract.tesseract_cmd = str(command_path)
            version = str(pytesseract.get_tesseract_version())
            available_languages = set(pytesseract.get_languages(config=""))
        except (pytesseract.TesseractNotFoundError, OSError):
            return PdfPageExtraction(page_number, "unavailable", None, None,
                "The local Tesseract runtime is unavailable.", None, None, None, ())
        languages = "+".join(language for language in ("eng", "deu") if language in available_languages)
        if not languages:
            return PdfPageExtraction(page_number, "unavailable", None, None,
                "Neither English nor German Tesseract language data is available.", None, version, None, ())
        try:
            from pymupdf import Matrix

            matrix = Matrix(OCR_RENDER_SCALE, OCR_RENDER_SCALE)
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)  # type: ignore[attr-defined]
            if pixmap.width * pixmap.height > MAX_OCR_PIXELS:
                raise DocumentImportError("OCR_PAGE_TOO_LARGE", "A PDF page exceeds the safe OCR rendering limit.")
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            data = pytesseract.image_to_data(
                image, lang=languages, config="--oem 1 --psm 6",
                output_type=pytesseract.Output.DICT, timeout=OCR_PAGE_TIMEOUT_SECONDS,
            )
        except RuntimeError as error:
            return PdfPageExtraction(page_number, "unavailable", None, None,
                "Local OCR timed out or could not process this page.", languages, version, None, ())
        except DocumentImportError:
            raise
        except Exception as error:
            return PdfPageExtraction(page_number, "unavailable", None, None,
                "Local OCR could not process this page.", languages, version, None, ())

        blocks: list[PdfTextBlock] = []
        confidences: list[float] = []
        for index, text in enumerate(data.get("text", [])):
            if not isinstance(text, str) or not text.strip():
                continue
            try:
                confidence = float(data["conf"][index])
            except (KeyError, TypeError, ValueError):
                confidence = None
            if confidence is not None and confidence >= 0:
                confidences.append(confidence)
            blocks.append(PdfTextBlock(
                len(blocks), text, float(data["left"][index]), float(data["top"][index]),
                float(data["left"][index] + data["width"][index]), float(data["top"][index] + data["height"][index]),
                confidence,
            ))
        raw_text = " ".join(block.text for block in blocks)
        warning = "OCR output may contain recognition or layout errors. Review against the source page."
        if not raw_text:
            warning = "OCR completed but no readable text was detected on this page."
        return PdfPageExtraction(page_number, "ocr", raw_text, self._clean_pdf_text(raw_text), warning,
            languages, version, (sum(confidences) / len(confidences)) if confidences else None, tuple(blocks))

    @staticmethod
    def _check_item_limit(items: Sequence[ExtractedItem]) -> None:
        if len(items) > MAX_EXTRACTED_ITEMS:
            raise DocumentImportError("CONTENT_TOO_LARGE", "The document exceeds the safe extraction limit.")

    @staticmethod
    def _column_letters(column_number: int) -> str:
        letters = ""
        while column_number:
            column_number, remainder = divmod(column_number - 1, 26)
            letters = chr(65 + remainder) + letters
        return letters


class DocumentImportService:
    """Coordinates vault import, deterministic parsing, and extraction persistence."""

    def __init__(self, storage: StorageService) -> None:
        if storage.database is None or storage.vault is None:
            raise DocumentImportError("STORAGE_UNAVAILABLE", "Local storage is not ready.")
        self._database = storage.database
        self._vault = storage.vault
        self._repository = DocumentMetadataRepository()
        self._parser = DeterministicDocumentParser()

    def import_authorized_paths(self, paths: object) -> Mapping[str, object]:
        if not isinstance(paths, list) or not (1 <= len(paths) <= MAX_IMPORT_FILES):
            raise DocumentImportError("INVALID_IMPORT", "Select between one and twenty documents.")
        results = []
        for value in paths:
            if not isinstance(value, str):
                results.append({"ok": False, "filename": "Unknown file", "error": self._error("INVALID_IMPORT", "Invalid selected file.")})
                continue
            results.append(self._import_one(Path(value)))
        return {"documents": results}

    def _import_one(self, source: Path) -> Mapping[str, object]:
        filename = source.name or "Selected file"
        try:
            suffix = source.suffix.lower()
            file_format = SUPPORTED_FORMATS.get(suffix)
            if file_format is None:
                raise DocumentImportError("UNSUPPORTED_FORMAT", "Supported formats are XLSX, XLS, CSV, and PDF.")
            try:
                size = source.lstat().st_size
            except OSError as error:
                raise DocumentImportError("UNREADABLE_FILE", "The selected file cannot be inspected.") from error
            if size > MAX_FILE_BYTES:
                raise DocumentImportError("FILE_TOO_LARGE", "Each document must be 50 MB or smaller.")
            vault_result = self._vault.import_file(source)
            document = vault_result.document
            if not vault_result.duplicate:
                self._set_document_format(document.id, file_format)
                self._parse_document(document)
            elif self._parse_retry_needed(document.id):
                # A duplicate is normally a no-op, but an earlier local parse
                # may have failed because of a corrected parser/runtime issue.
                # Re-selecting the same user-authorized file is an explicit,
                # safe retry of the immutable managed copy; it never creates
                # a second document or alters the original source file.
                self._set_document_format(document.id, file_format)
                self._parse_document(document)
            return {"ok": True, "document": self._summary(document.id), "duplicate": vault_result.duplicate}
        except (DocumentImportError, VaultImportError) as error:
            return {"ok": False, "filename": filename, "error": self._error(getattr(error, "code", "IMPORT_FAILED"), str(error))}

    def retry_parse(self, document_id: object) -> Mapping[str, object]:
        document = self._get_document(document_id)
        self._parse_document(document)
        return {"document": self._summary(document.id)}

    def _parse_retry_needed(self, document_id: str) -> bool:
        with self._database.connection() as connection:
            row = connection.execute(
                "SELECT processing_status FROM source_documents WHERE id = ?", (document_id,)
            ).fetchone()
        return row is not None and str(row["processing_status"]) in {"failed", "imported"}

    def list_documents(self) -> Mapping[str, object]:
        with self._database.connection() as connection:
            rows = connection.execute("SELECT id FROM source_documents ORDER BY imported_at DESC").fetchall()
        return {"documents": [self._summary(str(row["id"])) for row in rows]}

    def remove_document(self, document_id: object) -> Mapping[str, object]:
        """Remove one user-confirmed imported document and its local derived data.

        Confirmed domain records and their source evidence are never removed here.
        The user-facing operation only applies to an unused imported document,
        including its deterministic extractions and provisional AI results.
        """

        document = self._get_document(document_id)
        with self._database.transaction() as connection:
            active = connection.execute(
                "SELECT 1 FROM ai_processing_jobs WHERE source_document_id = ? AND status IN ('queued', 'running') LIMIT 1",
                (document.id,),
            ).fetchone()
            if active is not None:
                raise DocumentImportError("DOCUMENT_BUSY", "Wait for the active AI analysis to finish before removing this document.")
            referenced = any(
                connection.execute(f"SELECT 1 FROM {table} WHERE source_document_id = ? LIMIT 1", (document.id,)).fetchone()
                is not None
                for table in ("holding_positions", "claim_right_assignments", "evidence_references", "issues")
            )
            if referenced:
                raise DocumentImportError(
                    "DOCUMENT_IN_USE",
                    "This document is linked to confirmed records or an issue and cannot be removed here.",
                )
            connection.execute(
                "DELETE FROM ai_proposals WHERE job_id IN (SELECT id FROM ai_processing_jobs WHERE source_document_id = ?)",
                (document.id,),
            )
            connection.execute("DELETE FROM ai_processing_jobs WHERE source_document_id = ?", (document.id,))
            connection.execute(
                "DELETE FROM pdf_text_blocks WHERE pdf_page_extraction_id IN ("
                "SELECT p.id FROM pdf_page_extractions p JOIN document_extractions e ON e.id = p.extraction_id "
                "WHERE e.source_document_id = ?)",
                (document.id,),
            )
            connection.execute(
                "DELETE FROM pdf_page_extractions WHERE extraction_id IN ("
                "SELECT id FROM document_extractions WHERE source_document_id = ?)",
                (document.id,),
            )
            connection.execute(
                "DELETE FROM extracted_content WHERE extraction_id IN ("
                "SELECT id FROM document_extractions WHERE source_document_id = ?)",
                (document.id,),
            )
            connection.execute("DELETE FROM document_extractions WHERE source_document_id = ?", (document.id,))
            connection.execute("DELETE FROM source_documents WHERE id = ?", (document.id,))
        self._vault.remove_managed_copy(document)
        return {"documentId": document.id}

    def content(
        self, document_id: object, offset: object, limit: object, sheet_name: object = None,
        row_start: object = None, row_limit: object = None, column_start: object = None,
        column_limit: object = None, page_number: object = None,
    ) -> Mapping[str, object]:
        document = self._get_document(document_id)
        safe_offset = offset if isinstance(offset, int) and not isinstance(offset, bool) and offset >= 0 else 0
        safe_limit = limit if isinstance(limit, int) and not isinstance(limit, bool) and 1 <= limit <= 100 else 25
        safe_sheet_name = sheet_name if isinstance(sheet_name, str) and 0 < len(sheet_name) <= 255 else None
        safe_row_start = row_start if isinstance(row_start, int) and not isinstance(row_start, bool) and 1 <= row_start <= 1_000_000 else 1
        safe_row_limit = row_limit if isinstance(row_limit, int) and not isinstance(row_limit, bool) and 1 <= row_limit <= MAX_GRID_ROWS else MAX_GRID_ROWS
        safe_column_start = column_start if isinstance(column_start, int) and not isinstance(column_start, bool) and 1 <= column_start <= 16_384 else 1
        safe_column_limit = column_limit if isinstance(column_limit, int) and not isinstance(column_limit, bool) and 1 <= column_limit <= MAX_GRID_COLUMNS else MAX_GRID_COLUMNS
        safe_page_number = page_number if isinstance(page_number, int) and not isinstance(page_number, bool) and page_number >= 1 else None
        with self._database.connection() as connection:
            extraction = connection.execute("""SELECT id, parser_version, status FROM document_extractions
                WHERE source_document_id = ? ORDER BY started_at DESC LIMIT 1""", (document.id,)).fetchone()
            if extraction is None:
                return {"document": self._summary(document.id), "references": [], "hasMore": False}
            file_format = self._format_for_document(connection, document.id)
            if file_format in {"xlsx", "xls", "csv"}:
                return self._spreadsheet_content_response(
                    connection, document.id, extraction, file_format, safe_sheet_name, safe_row_start, safe_row_limit,
                    safe_column_start, safe_column_limit,
                )
            if file_format == "pdf":
                return self._pdf_content_response(connection, document.id, extraction, safe_page_number)
            rows = connection.execute("""SELECT content_kind, sheet_name, row_number, column_number, page_number,
                value_text, value_type, formula_text, is_empty, source_reference FROM extracted_content
                WHERE extraction_id = ? ORDER BY page_number, row_number, column_number, source_reference LIMIT ? OFFSET ?""", (extraction["id"], safe_limit + 1, safe_offset)).fetchall()
        summary = self._summary(document.id)
        references: list[dict[str, object]] = []
        has_more = len(rows) > safe_limit
        for row in rows[:safe_limit]:
            reference = dict(row) | {"is_empty": bool(row["is_empty"])}
            value = reference.get("value_text")
            if isinstance(value, str) and len(value) > MAX_PREVIEW_VALUE_CHARS:
                reference["value_text"] = value[:MAX_PREVIEW_VALUE_CHARS] + "…"
                reference["value_truncated"] = True
            candidate = {"document": summary, "parserVersion": extraction["parser_version"],
                "references": [*references, reference], "hasMore": True}
            if len(json.dumps(candidate, ensure_ascii=False, separators=(",", ":")).encode("utf-8")) > MAX_PREVIEW_RESPONSE_BYTES:
                has_more = True
                break
            references.append(reference)
        return {"document": summary, "parserVersion": extraction["parser_version"],
            "references": references, "hasMore": has_more}

    def _spreadsheet_content_response(
        self, connection: sqlite3.Connection, document_id: str, extraction: sqlite3.Row, file_format: str,
        requested_sheet: str | None, row_start: int, row_limit: int, column_start: int, column_limit: int,
    ) -> Mapping[str, object]:
        worksheet_rows = connection.execute("""SELECT COALESCE(sheet_name, 'CSV') AS sheet_name,
            MIN(worksheet_index) AS worksheet_index, MAX(row_number) AS max_row, MAX(column_number) AS max_column,
            COUNT(*) AS cell_count, SUM(CASE WHEN is_empty = 0 THEN 1 ELSE 0 END) AS non_empty_cell_count
            FROM extracted_content WHERE extraction_id = ? AND content_kind IN ('spreadsheet_cell', 'csv_cell')
            GROUP BY COALESCE(sheet_name, 'CSV')
            ORDER BY CASE WHEN MIN(worksheet_index) IS NULL THEN 1 ELSE 0 END, MIN(worksheet_index), sheet_name COLLATE BINARY""",
            (extraction["id"],)).fetchall()
        worksheets = [{
            "name": str(row["sheet_name"]), "worksheetIndex": row["worksheet_index"],
            "maxRow": int(row["max_row"] or 0), "maxColumn": int(row["max_column"] or 0),
            "cellCount": int(row["cell_count"]), "nonEmptyCellCount": int(row["non_empty_cell_count"] or 0),
        } for row in worksheet_rows]
        selected_sheet = requested_sheet if requested_sheet in {sheet["name"] for sheet in worksheets} else (worksheets[0]["name"] if worksheets else None)
        cells: list[dict[str, object]] = []
        preview_truncated = False
        selected_summary = next((sheet for sheet in worksheets if sheet["name"] == selected_sheet), None)
        if selected_sheet is not None:
            rows = connection.execute("""SELECT row_number, column_number, value_text, value_type, formula_text,
                is_empty, source_reference FROM extracted_content
                WHERE extraction_id = ? AND content_kind IN ('spreadsheet_cell', 'csv_cell')
                AND COALESCE(sheet_name, 'CSV') = ? AND row_number >= ? AND row_number < ?
                AND column_number >= ? AND column_number < ?
                AND (is_empty = 0 OR formula_text IS NOT NULL)
                ORDER BY row_number ASC, column_number ASC, source_reference ASC""",
                (extraction["id"], selected_sheet, row_start, row_start + row_limit, column_start, column_start + column_limit)).fetchall()
            for row in rows:
                cell = dict(row) | {"is_empty": bool(row["is_empty"])}
                value = cell.get("value_text")
                if isinstance(value, str) and len(value) > MAX_GRID_VALUE_CHARS:
                    cell["value_text"] = value[:MAX_GRID_VALUE_CHARS] + "…"
                    cell["value_truncated"] = True
                formula = cell.get("formula_text")
                if isinstance(formula, str) and len(formula) > MAX_GRID_VALUE_CHARS:
                    cell["formula_text"] = formula[:MAX_GRID_VALUE_CHARS] + "…"
                    cell["formula_truncated"] = True
                candidate_cells = [*cells, cell]
                if len(json.dumps(candidate_cells, ensure_ascii=False, separators=(",", ":")).encode("utf-8")) > MAX_GRID_RESPONSE_BYTES:
                    preview_truncated = True
                    break
                cells.append(cell)
        return {
            "document": self._summary(document_id), "parserVersion": extraction["parser_version"], "contentType": "spreadsheet",
            "extraction": self._extraction_summary(connection, extraction), "worksheets": worksheets,
            "grid": {
                "sheetName": selected_sheet, "rowStart": row_start, "rowLimit": row_limit,
                "maxRow": selected_summary["maxRow"] if selected_summary else 0,
                "columnStart": column_start, "columnLimit": column_limit,
                "maxColumn": int(selected_summary["maxColumn"]) if selected_summary else 0,
                "cells": cells, "emptyCellsOmitted": True, "previewTruncated": preview_truncated,
            }, "references": [], "hasMore": bool(selected_summary and row_start + row_limit <= selected_summary["maxRow"]),
        }

    @staticmethod
    def _format_for_document(connection: sqlite3.Connection, document_id: str) -> str:
        row = connection.execute("SELECT file_format FROM source_documents WHERE id = ?", (document_id,)).fetchone()
        return str(row["file_format"]) if row is not None else "unknown"

    def _pdf_content_response(
        self, connection: sqlite3.Connection, document_id: str, extraction: sqlite3.Row, requested_page: int | None
    ) -> Mapping[str, object]:
        pages = connection.execute("SELECT page_number FROM pdf_page_extractions WHERE extraction_id = ? ORDER BY page_number", (extraction["id"],)).fetchall()
        page_numbers = [int(page["page_number"]) for page in pages]
        selected_page = requested_page if requested_page in page_numbers else (page_numbers[0] if page_numbers else None)
        previous_page = max((page for page in page_numbers if selected_page is not None and page < selected_page), default=None)
        next_page = min((page for page in page_numbers if selected_page is not None and page > selected_page), default=None)
        rows = connection.execute("""SELECT id, page_number, extraction_method, cleaned_text, raw_text, quality_warning,
            ocr_language, ocr_engine_version, mean_confidence, block_count FROM pdf_page_extractions
            WHERE extraction_id = ? AND page_number = ?""", (extraction["id"], selected_page)).fetchall() if selected_page is not None else []
        references: list[dict[str, object]] = []
        for row in rows:
            cleaned_text = str(row["cleaned_text"] or "")
            raw_text = str(row["raw_text"] or "")
            block_rows = connection.execute("""SELECT block_index, x0, y0, x1, y1, text, confidence,
                is_table_candidate FROM pdf_text_blocks WHERE pdf_page_extraction_id = ? ORDER BY block_index LIMIT 50""",
                (row["id"],)).fetchall()
            reference: dict[str, object] = {
                "content_kind": "pdf_page",
                "source_reference": f"page {row['page_number']}",
                "page_number": int(row["page_number"]),
                "value_text": cleaned_text[:MAX_PREVIEW_VALUE_CHARS],
                "raw_text": raw_text[:MAX_PREVIEW_VALUE_CHARS],
                "value_type": row["extraction_method"],
                "extraction_method": row["extraction_method"],
                "quality_warning": row["quality_warning"],
                "ocr_language": row["ocr_language"],
                "ocr_engine_version": row["ocr_engine_version"],
                "mean_confidence": row["mean_confidence"],
                "block_count": int(row["block_count"]),
                "is_empty": not bool(cleaned_text.strip()),
                "blocks": [dict(block) | {"is_table_candidate": bool(block["is_table_candidate"])} for block in block_rows],
            }
            if len(cleaned_text) > MAX_PREVIEW_VALUE_CHARS:
                reference["value_truncated"] = True
            references.append(reference)
        return {"document": self._summary(document_id), "parserVersion": extraction["parser_version"], "contentType": "pdf",
            "extraction": self._extraction_summary(connection, extraction), "references": references,
            "pagination": {"pageNumber": selected_page, "pageCount": len(page_numbers),
                "hasPrevious": previous_page is not None, "hasNext": next_page is not None,
                "previousPageNumber": previous_page, "nextPageNumber": next_page},
            "hasMore": next_page is not None}

    @staticmethod
    def _extraction_summary(connection: sqlite3.Connection, extraction: sqlite3.Row) -> Mapping[str, object]:
        items = connection.execute("""SELECT COUNT(*) AS item_count,
            SUM(CASE WHEN is_empty = 0 THEN 1 ELSE 0 END) AS non_empty_count,
            MAX(page_number) AS page_count FROM extracted_content WHERE extraction_id = ?""", (extraction["id"],)).fetchone()
        pdf = connection.execute("""SELECT COUNT(*) AS page_count, SUM(block_count) AS block_count,
            SUM(CASE WHEN extraction_method = 'ocr' THEN 1 ELSE 0 END) AS ocr_pages,
            SUM(CASE WHEN extraction_method = 'native_text' THEN 1 ELSE 0 END) AS native_pages,
            SUM(CASE WHEN quality_warning IS NOT NULL THEN 1 ELSE 0 END) AS warning_count
            FROM pdf_page_extractions WHERE extraction_id = ?""", (extraction["id"],)).fetchone()
        return {
            "status": extraction["status"], "totalItems": int(items["item_count"]),
            "nonEmptyItems": int(items["non_empty_count"] or 0), "pageCount": int(pdf["page_count"] or items["page_count"] or 0),
            "textBlockCount": int(pdf["block_count"] or 0), "nativePageCount": int(pdf["native_pages"] or 0),
            "ocrPageCount": int(pdf["ocr_pages"] or 0), "warningCount": int(pdf["warning_count"] or 0),
            "truncated": False,
        }

    def _parse_document(self, document: DocumentRecord) -> None:
        with self._database.connection() as connection:
            row = connection.execute("SELECT file_format FROM source_documents WHERE id = ?", (document.id,)).fetchone()
        file_format = str(row["file_format"]) if row is not None else "unknown"
        extraction_id, started_at = str(uuid.uuid4()), _now()
        with self._database.transaction() as connection:
            connection.execute("""INSERT INTO document_extractions (id, source_document_id, parser_version, status, started_at)
                VALUES (?, ?, ?, 'processing', ?)""", (extraction_id, document.id, PARSER_VERSION, started_at))
        try:
            result = self._parser.parse(self._vault.managed_path_for(document), file_format)
            with self._database.transaction() as connection:
                connection.executemany("""INSERT INTO extracted_content (id, extraction_id, content_kind, sheet_name, row_number,
                    column_number, page_number, worksheet_index, value_text, value_type, formula_text, is_empty, source_reference)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", [
                    (str(uuid.uuid4()), extraction_id, item.content_kind, item.sheet_name, item.row_number, item.column_number,
                     item.page_number, item.worksheet_index, item.value_text, item.value_type, item.formula_text, int(item.is_empty), item.source_reference)
                    for item in result.items
                ])
                self._insert_pdf_pages(connection, extraction_id, result.pdf_pages)
                connection.execute("UPDATE document_extractions SET status = 'parsed', completed_at = ? WHERE id = ?", (_now(), extraction_id))
                connection.execute("""UPDATE source_documents SET processing_status = 'parsed', parser_version = ?, parsed_at = ?, processing_error = NULL
                    WHERE id = ?""", (PARSER_VERSION, _now(), document.id))
        except ExtractionUnsupported as error:
            self._mark_parse_failure(document.id, extraction_id, "unsupported", error)
        except (DocumentImportError, VaultStorageError) as error:
            self._mark_parse_failure(document.id, extraction_id, "failed", error)
        except Exception as error:
            # Never leave a failed attempt in the transient processing state.
            self._mark_parse_failure(
                document.id,
                extraction_id,
                "failed",
                DocumentImportError("PARSE_FAILED", "The document could not be parsed safely."),
            )

    @staticmethod
    def _insert_pdf_pages(
        connection: sqlite3.Connection, extraction_id: str, pages: Sequence[PdfPageExtraction]
    ) -> None:
        for page in pages:
            page_id = str(uuid.uuid4())
            connection.execute("""INSERT INTO pdf_page_extractions (
                id, extraction_id, page_number, extraction_method, raw_text, cleaned_text, quality_warning,
                ocr_language, ocr_engine_version, mean_confidence, block_count
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", (
                page_id, extraction_id, page.page_number, page.extraction_method, page.raw_text, page.cleaned_text,
                page.quality_warning, page.ocr_language, page.ocr_engine_version, page.mean_confidence, len(page.blocks),
            ))
            connection.executemany("""INSERT INTO pdf_text_blocks (
                id, pdf_page_extraction_id, block_index, x0, y0, x1, y1, text, confidence, is_table_candidate
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", [
                (str(uuid.uuid4()), page_id, block.block_index, block.x0, block.y0, block.x1, block.y1,
                 block.text, block.confidence, int(block.is_table_candidate))
                for block in page.blocks
            ])

    def _mark_parse_failure(self, document_id: str, extraction_id: str, status: str, error: DocumentImportError) -> None:
        with self._database.transaction() as connection:
            connection.execute("""UPDATE document_extractions SET status = ?, completed_at = ?, error_code = ?, error_message = ? WHERE id = ?""",
                (status, _now(), error.code, str(error), extraction_id))
            connection.execute("""UPDATE source_documents SET processing_status = ?, parser_version = ?, parsed_at = ?, processing_error = ? WHERE id = ?""",
                (status, PARSER_VERSION, _now(), str(error), document_id))

    def _set_document_format(self, document_id: str, file_format: str) -> None:
        with self._database.transaction() as connection:
            connection.execute("UPDATE source_documents SET file_format = ? WHERE id = ?", (file_format, document_id))

    def _get_document(self, document_id: object) -> DocumentRecord:
        if not isinstance(document_id, str) or not document_id:
            raise DocumentImportError("INVALID_DOCUMENT", "The document reference is invalid.")
        with self._database.connection() as connection:
            row = connection.execute("""SELECT id, original_filename, managed_relative_path, content_sha256, file_size_bytes, imported_at
                FROM source_documents WHERE id = ?""", (document_id,)).fetchone()
        if row is None:
            raise DocumentImportError("DOCUMENT_NOT_FOUND", "The selected document is unavailable.")
        return DocumentMetadataRepository._from_row(row)

    def _summary(self, document_id: str) -> Mapping[str, object]:
        with self._database.connection() as connection:
            row = connection.execute("""SELECT d.id, d.original_filename, d.file_size_bytes, d.imported_at, d.file_format,
                d.processing_status, d.processing_error,
                CASE WHEN EXISTS (
                    SELECT 1 FROM ai_processing_jobs j
                    WHERE j.source_document_id = d.id
                      AND j.status = 'completed'
                      AND j.extraction_id = (
                        SELECT e.id FROM document_extractions e
                        WHERE e.source_document_id = d.id AND e.status = 'parsed'
                        ORDER BY e.started_at DESC, e.id DESC LIMIT 1
                      )
                ) THEN 'reviewed' ELSE 'new' END AS ai_review_status
                FROM source_documents d WHERE d.id = ?""", (document_id,)).fetchone()
        if row is None:
            raise DocumentImportError("DOCUMENT_NOT_FOUND", "The selected document is unavailable.")
        return dict(row)

    @staticmethod
    def _error(code: str, message: str) -> Mapping[str, str]:
        return {"code": code, "message": message}
