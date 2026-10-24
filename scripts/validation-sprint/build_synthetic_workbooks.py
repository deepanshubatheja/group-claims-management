"""Build the fictional XLSX sources and independently specified ground truth.

This generator has no Node or Codex-only dependency.  It deliberately uses
only the backend's declared ``openpyxl`` dependency so a clean development
checkout can recreate the synthetic validation scenario.
"""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "EEF5FB"
AMBER = "FFF2CC"


def _date(value: str) -> datetime:
    return datetime.fromisoformat(f"{value}T00:00:00")


def _title(sheet: Any, text: str, end_column: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    cell = sheet.cell(1, 1, text)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(bold=True, color="FFFFFF", size=14)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26


def _table(sheet: Any, title: str, headings: list[str], rows: list[list[object]]) -> None:
    _title(sheet, title, len(headings))
    for column, heading in enumerate(headings, 1):
        cell = sheet.cell(3, column, heading)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True, color=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row_number, row in enumerate(rows, 4):
        for column, value in enumerate(row, 1):
            sheet.cell(row_number, column, value).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A4"


def _holdings_workbook(path: Path, statement_date: str, issue_date: str, amounts: list[int], notes: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Holdings"
    _title(sheet, "Synthetic holdings statement", 10)
    metadata = (("Statement date", _date(statement_date)), ("Issue date", _date(issue_date)), ("Investor", "Redwood Investors LP"))
    for row_number, (label, value) in enumerate(metadata, 3):
        sheet.cell(row_number, 1, label)
        sheet.cell(row_number, 2, value)
    for cell in (sheet["B3"], sheet["B4"]):
        cell.number_format = "yyyy-mm-dd"
    headings = ["Account", "Custodian", "Subfund", "Security ISIN", "Security", "Currency", "Nominal amount", "As-of date", "Source status", "Notes"]
    for column, heading in enumerate(headings, 1):
        cell = sheet.cell(7, column, heading)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True, color=NAVY)
    rows = [
        ["CUST-ALPHA-01", "Atlas Custody Bank", "Horizon Credit Fund A (SF-A)", "XS0123456789", "Example Credit Notes 6.25% Perpetual", "USD", amounts[0], _date(statement_date), "Confirmed", notes[0]],
        ["CUST-ALPHA-02", "Atlas Custody Bank", "Horizon Credit Fund B (SF-B)", "XS0123456789", "Example Credit Notes 6.25% Perpetual", "USD", amounts[1], _date(statement_date), "Confirmed", notes[1]],
        ["CUST-ALPHA-03", "Northbridge Securities Custody", "Meridian Income Fund C (SF-C)", "XS0123456789", "Example Credit Notes 6.25% Perpetual", "USD", amounts[2], _date(statement_date), "Confirmed", notes[2]],
    ]
    for row_number, row in enumerate(rows, 8):
        for column, value in enumerate(row, 1):
            sheet.cell(row_number, column, value).alignment = Alignment(wrap_text=True, vertical="top")
    for row_number in range(8, 11):
        sheet.cell(row_number, 7).number_format = "#,##0"
        sheet.cell(row_number, 8).number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A8"
    workbook.save(path)


def _actual_values(output_dir: Path) -> dict[str, object]:
    report_path = output_dir / "isolated_validation_report.json"
    if not report_path.exists():
        return {"imports": 0, "hashes": "Pending", "holdings": 0, "assignments": 0, "issues": 0, "billing": 0, "duplicate": "Pending"}
    report = json.loads(report_path.read_text(encoding="utf-8"))
    return {
        "imports": report["imports"]["actual"],
        "hashes": "Yes" if report["vaultHashes"]["actual"] else "Pending",
        "holdings": report["domainPersistence"]["actual"]["holdings"],
        "assignments": report["domainPersistence"]["actual"]["claimAssignments"],
        "issues": report["domainPersistence"]["actual"]["issues"],
        "billing": 120000 if report["billing"]["pass"] else 0,
        "duplicate": "Yes" if report["duplicateDetection"]["actual"] else "Pending",
    }


def _ground_truth_workbook(path: Path, actual: dict[str, object]) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _title(summary, "ClaimFlow synthetic validation — independent ground truth", 6)
    summary.merge_cells("A3:F3")
    summary["A3"] = "Scenario ground truth is specified independently from ClaimFlow calculations and AI output. Actual values are populated after the isolated run."
    summary["A3"].fill = PatternFill("solid", fgColor=PALE)
    headings = ["Control", "Expected", "Actual at run", "Comparison", "Status", "Notes"]
    for column, heading in enumerate(headings, 1):
        summary.cell(5, column, heading).fill = PatternFill("solid", fgColor=BLUE)
        summary.cell(5, column).font = Font(bold=True, color=NAVY)
    summary_rows = [
        ["Imported source documents", 6, actual["imports"], "=C6-B6", '=IF(C6=B6,"PASS","FAIL")', "Independent source-count control"],
        ["Original hashes equal vault copies", "Yes", actual["hashes"], None, '=IF(C7=B7,"PASS","FAIL")', "SHA-256 validation"],
        ["Dated holding records persisted", 6, actual["holdings"], "=C8-B8", '=IF(C8=B8,"PASS","FAIL")', "Expected from a complete reconciliation workflow"],
        ["Securities transfer events persisted", 1, 0, "=C9-B9", '=IF(C9=B9,"PASS","FAIL")', "Separate from a claim-right assignment"],
        ["Claim-right assignments persisted", 1, actual["assignments"], "=C10-B10", '=IF(C10=B10,"PASS","FAIL")', "Expected legal ownership event"],
        ["Contradiction issues persisted", 1, actual["issues"], "=C11-B11", '=IF(C11=B11,"PASS","FAIL")', "Never silently choose a conflicting balance"],
        ["Pro-rata allocation total", 120000, actual["billing"], "=C12-B12", '=IF(C12=B12,"PASS","FAIL")', "Explicit instruction only"],
        ["Live cloud-AI test", "Not run", "Not run", None, "Not run", "Requires credential and explicit authorization"],
    ]
    for row_number, row in enumerate(summary_rows, 6):
        for column, value in enumerate(row, 1):
            summary.cell(row_number, column, value)
    summary.merge_cells("A15:F15")
    summary["A15"] = "Expected allocation: Fund A 400,000 / 1,400,000 × USD 120,000 = USD 34,285.71; Fund B 1,000,000 / 1,400,000 × USD 120,000 = USD 85,714.29; Fund C is excluded."
    summary["A15"].fill = PatternFill("solid", fgColor=AMBER)

    documents = workbook.create_sheet("Documents")
    _table(documents, "Documents", ["Source document", "Format", "Issue date", "Effective/as-of date", "Expected role", "Expected reference"], [
        ["01_opening_holdings_2025-01-31.xlsx", "XLSX", _date("2025-02-02"), _date("2025-01-31"), "Opening dated holdings", "Holdings!A8:J10"],
        ["02_partial_transfer_effective_2025-02-15.pdf", "PDF", _date("2025-02-20"), _date("2025-02-15"), "Partial securities transfer", "page 1"],
        ["03_claim_right_assignment_effective_2025-02-18.pdf", "PDF", _date("2025-02-21"), _date("2025-02-18"), "Claim-right assignment only", "page 1"],
        ["04_later_holdings_2025-02-28.xlsx", "XLSX", _date("2025-03-03"), _date("2025-02-28"), "Later dated holdings", "Holdings!A8:J10"],
        ["05_contradictory_custodian_statement_2025-02-28.pdf", "PDF", _date("2025-03-10"), _date("2025-02-28"), "Contradictory B holding", "page 1"],
        ["06_approved_billing_instruction_2025-03-01.pdf", "PDF", _date("2025-03-01"), _date("2025-02-01"), "Explicit billing instruction", "page 1"],
    ])
    holdings = workbook.create_sheet("Expected Holdings")
    _table(holdings, "Expected Holdings", ["As-of date", "Subfund", "Account", "ISIN", "Nominal amount", "Source", "Reference"], [
        [_date("2025-01-31"), "Horizon Credit Fund A (SF-A)", "CUST-ALPHA-01", "XS0123456789", 600000, "01_opening_holdings_2025-01-31.xlsx", "Holdings!A8:J8"],
        [_date("2025-01-31"), "Horizon Credit Fund B (SF-B)", "CUST-ALPHA-02", "XS0123456789", 800000, "01_opening_holdings_2025-01-31.xlsx", "Holdings!A9:J9"],
        [_date("2025-01-31"), "Meridian Income Fund C (SF-C)", "CUST-ALPHA-03", "XS0123456789", 500000, "01_opening_holdings_2025-01-31.xlsx", "Holdings!A10:J10"],
        [_date("2025-02-28"), "Horizon Credit Fund A (SF-A)", "CUST-ALPHA-01", "XS0123456789", 400000, "04_later_holdings_2025-02-28.xlsx", "Holdings!A8:J8"],
        [_date("2025-02-28"), "Horizon Credit Fund B (SF-B)", "CUST-ALPHA-02", "XS0123456789", 1000000, "04_later_holdings_2025-02-28.xlsx", "Holdings!A9:J9"],
        [_date("2025-02-28"), "Meridian Income Fund C (SF-C)", "CUST-ALPHA-03", "XS0123456789", 500000, "04_later_holdings_2025-02-28.xlsx", "Holdings!A10:J10"],
    ])
    events = workbook.create_sheet("Expected Events")
    _table(events, "Expected Events", ["Effective date", "Event type", "From", "To", "ISIN / claim", "Quantity / percent", "Source", "Reference"], [
        [_date("2025-02-15"), "Partial securities transfer", "Horizon Credit Fund A (SF-A)", "Horizon Credit Fund B (SF-B)", "XS0123456789", "200,000 nominal", "02_partial_transfer_effective_2025-02-15.pdf", "page 1"],
        [_date("2025-02-18"), "Claim-right assignment", "Redwood Investors LP", "Northstar Claims SPV", "CR-2025-ALPHA", "100%", "03_claim_right_assignment_effective_2025-02-18.pdf", "page 1"],
    ])
    discrepancies = workbook.create_sheet("Expected Discrepancies")
    _table(discrepancies, "Expected Discrepancies", ["As-of date", "Subject", "Expected discrepancy", "Source A", "Source B", "Required treatment"], [[_date("2025-02-28"), "Horizon Credit Fund B (SF-B)", "900,000 nominal conflicts with 1,000,000 nominal", "04_later_holdings_2025-02-28.xlsx, Holdings!A9:J9", "05_contradictory_custodian_statement_2025-02-28.pdf, page 1", "Record unresolved issue; do not silently choose a balance."]])
    billing = workbook.create_sheet("Billing Ground Truth")
    _table(billing, "Billing Ground Truth", ["Participant", "Eligible nominal", "Formula", "Percentage", "Allocation (USD)", "Source / status"], [
        ["Horizon Credit Fund A (SF-A)", 400000, "=B4/$B$6", "=C4", "=ROUND(C4*$B$2,2)", "Eligible — approved billing instruction"],
        ["Horizon Credit Fund B (SF-B)", 1000000, "=B5/$B$6", "=C5", "=ROUND(C5*$B$2,2)", "Eligible — approved billing instruction"],
        ["Total eligible nominal", "=SUM(B4:B5)", None, "=SUM(D4:D5)", "=SUM(E4:E5)", "Must reconcile to fee pool"],
        ["Meridian Income Fund C (SF-C)", 500000, "Excluded", 0, 0, "Excluded — not participating"],
    ])
    billing["A2"] = "Fee pool (USD)"
    billing["B2"] = 120000
    billing["A2"].fill = billing["B2"].fill = PatternFill("solid", fgColor=AMBER)
    actual_sheet = workbook.create_sheet("Actual vs Expected")
    _table(actual_sheet, "Actual vs Expected", ["Test / control", "Expected", "Actual isolated run", "Difference", "Pass / fail", "Evidence"], [
        ["Documents imported and parsed", 6, actual["imports"], "=C4-B4", '=IF(C4=B4,"PASS","FAIL")', "Import service response"],
        ["Duplicate import prevented", "Yes", actual["duplicate"], None, '=IF(C5=B5,"PASS","FAIL")', "SHA-256 duplicate detection"],
        ["Vault hash equality", "Yes", actual["hashes"], None, '=IF(C6=B6,"PASS","FAIL")', "Original and managed file SHA-256"],
        ["Dated holding records persisted", 6, actual["holdings"], "=C7-B7", '=IF(C7=B7,"PASS","FAIL")', "holding_positions table"],
        ["Transfer events persisted", 1, 0, "=C8-B8", '=IF(C8=B8,"PASS","FAIL")', "No dedicated securities transfer event table"],
        ["Claim assignments persisted", 1, actual["assignments"], "=C9-B9", '=IF(C9=B9,"PASS","FAIL")', "claim_right_assignments table"],
        ["Contradiction issue persisted", 1, actual["issues"], "=C10-B10", '=IF(C10=B10,"PASS","FAIL")', "issues table"],
        ["Billing allocation total", 120000, actual["billing"], "=C11-B11", '=IF(C11=B11,"PASS","FAIL")', "MatterReviewService instruction calculation"],
    ])
    workbook.save(path)


def main() -> None:
    output_dir = Path(sys.argv[1]).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    _holdings_workbook(output_dir / "01_opening_holdings_2025-01-31.xlsx", "2025-01-31", "2025-02-02", [600000, 800000, 500000], ["Opening position", "Opening position", "Opening position"])
    _holdings_workbook(output_dir / "04_later_holdings_2025-02-28.xlsx", "2025-02-28", "2025-03-03", [400000, 1000000, 500000], ["After 200,000 partial transfer", "After 200,000 partial transfer", "No change"])
    _ground_truth_workbook(output_dir / "ClaimFlow_Synthetic_Ground_Truth.xlsx", _actual_values(output_dir))


if __name__ == "__main__":
    main()
